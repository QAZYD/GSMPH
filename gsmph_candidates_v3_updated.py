"""
GSMPH localization-candidate extractor (v3)

What changed from v2 (driven by real failures in the sample output)
  1. Template stripping. The CSV 'text' column holds the full training string
     "<s>[INST]question [/INST] solution #### n </s>". v2 fed all of that to the
     NER models, which produced spans like 's>[INST]Jerry' and duplicate hits from
     the solution text. v3 extracts only the question.
  2. Food gating. FoodBaseBERT fires on animals, clothing, flowers and colors
     ('birds', 'starfish', 'shoes', 'violet'). A food hit now survives as Tier 2 only
     if WordNet agrees; otherwise it is demoted to "Food?" (Tier 3). Generic words
     ('food', 'meal') are dropped. Threshold raised to 0.90.
  3. Cross-category name fix. A single-word ORG/Place guess is re-labeled
     Proper Name when the word is a known first name or is the subject of a verb.
  4. reject_terms.txt. Terms you reject repeatedly are hidden on the next run
     (build it with --learn-rejects).
  5. Score column (model confidence) so you can pick thresholds from your own data.
  6. Object support votes. A WordNet-only object stays Tier 3; it moves to Tier 2 when another INDEPENDENT
     source agrees: 'counted' (the problem counts it: "7 starfish", "how many gnomes", "a box of") or 'concrete'
     (optional Brysbaert concreteness ratings, --concreteness FILE). A WordNet hypernym walk fills the optional
     Suggested Pool column (same source as WordNet, so NOT a vote). v3 never imports extract_pool.py.
  7. --precision-report is case-insensitive and says how many rows it read/counted.

Usage
  python gsmph_candidates_v3.py --csv gsm8k.csv --limit 100 --out sample.xlsx
  (review: set the Decision column to Keep / Reject; save under a NEW name, close Excel)
  python gsmph_candidates_v3.py --precision-report reviewed.xlsx
  python gsmph_candidates_v3.py --learn-rejects reviewed.xlsx
"""
import argparse
import csv
import os
import re
from collections import defaultdict
from dataclasses import dataclass
from functools import lru_cache
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

try:
    from transformers import pipeline as hf_pipeline
except ImportError:
    hf_pipeline = None

try:
    import nltk
    from nltk.corpus import names as nltk_names, wordnet as wn
except ImportError:
    nltk = nltk_names = wn = None


# ============================================================================
# CONFIG
# ============================================================================
SPACY_MODEL = "en_core_web_lg"           # falls back to en_core_web_sm
HF_NER_MODEL = "dslim/bert-base-NER"     # PER / LOC / ORG / MISC (MISC ignored)
FOOD_MODEL = "Dizex/FoodBaseBERT-NER"
NER_MIN_SCORE = 0.85
FOOD_MIN_SCORE = 0.90
WN_SENSES = 2                            # WordNet senses consulted for generic objects
HF_BATCH = 32
GSM8K_CSV_PATH = "gsm8k.csv"
QUESTION_COLUMN_CANDIDATES = ("question", "problem", "text")
REJECT_FILE = "reject_terms.txt"
CONCRETE_MIN = 4.0                       # Brysbaert mean concreteness (1-5); >= this supports "physical thing"
CONCRETENESS = {}                        # word -> rating, filled by --concreteness
EXTRA_NAMES_FILE = "extra_names.txt"      # one first name per line; added to the name gazetteer
REJECT_TERMS = set()                     # filled by extract(); lowercase


# ============================================================================
# WORD LISTS  (extend these as you review more of the dataset)
# ============================================================================
MONTHS = {"january", "february", "march", "april", "may", "june", "july",
          "august", "september", "october", "november", "december"}
DAYS = {"monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"}

KINSHIP = {"friend", "parent", "grandparent", "brother", "sister", "mother", "father",
           "mom", "dad", "mommy", "daddy", "son", "daughter", "cousin", "aunt", "uncle",
           "husband", "wife", "family", "kid", "child", "children", "classmate",
           "neighbor", "neighbour", "people", "person", "man", "woman", "boy", "girl"}
ABSTRACT = {"half", "quarter", "third", "purpose", "amount", "quantity", "total", "sum",
            "difference", "remainder", "percentage", "percent", "rate", "value", "number",
            "average", "ratio", "portion", "fraction"}
TIME_WORDS = {"yesterday", "today", "tomorrow", "hour", "minute", "day", "week", "month",
              "year", "shift", "overtime", "time", "second", "weekend", "morning", "night"}
GENERIC = {"thing", "item", "stuff", "piece", "type", "kind", "set", "pair", "lot",
           "bunch", "row", "side", "part", "end", "way", "rest", "dollar", "cent"}
GENERIC_FOOD = {"food", "foods", "meal", "meals", "diet", "drink", "drinks", "water", "dinner", "lunch", "breakfast",
                "supper", "brunch", "ingredient", "ingredients", "goodie", "goodies", "leftover", "leftovers", "treat", "treats"}
NOT_OBJECTS = {"work", "weight", "round", "shopping", "internet", "snowflake", "snow", "rain", "weather", "money",
               "price", "cost", "game", "trip", "ride", "job", "class", "lesson", "test", "week"}
EXCLUDED_NOUNS = KINSHIP | ABSTRACT | TIME_WORDS | GENERIC | GENERIC_FOOD | NOT_OBJECTS

# Lexicons: category -> (terms, tier). Case-insensitive, optional plural 's'.
LEXICONS = {
    "Event": ([
        "Christmas", "Christmas Eve", "Halloween", "Thanksgiving", "Easter",
        "Valentine's Day", "Valentine's", "New Year's Eve", "New Year's Day", "New Year",
        "Independence Day", "Fourth of July", "July 4th", "Memorial Day", "Labor Day",
        "Black Friday", "Cyber Monday", "Hanukkah", "Super Bowl", "Mother's Day",
        "Father's Day", "St. Patrick's Day", "Groundhog Day", "trick-or-treat",
        "trick or treat", "Easter egg hunt"], 2),
    "Culture": ([
        "bake sale", "garage sale", "yard sale", "lemonade stand", "prom", "homecoming",
        "spelling bee", "science fair", "talent show", "PTA", "varsity", "freshman",
        "sophomore", "kindergarten", "recess", "cafeteria", "locker", "baseball",
        "softball", "football", "hockey", "touchdown", "home run", "quarterback",
        "slumber party", "sleepover", "boy scouts", "girl scouts", "summer camp",
        "road trip", "babysitting", "babysitter"], 2),
    "Brand": ([
        "Walmart", "Costco", "McDonald's", "Starbucks", "Amazon", "eBay", "Netflix",
        "Disney", "Lego", "iPhone", "Xbox", "PlayStation", "Nintendo", "Barbie",
        "Pokemon", "Pokémon", "Subway", "Home Depot", "Whole Foods", "Walgreens", "CVS",
        "Kroger", "Trader Joe's", "Dunkin", "Pizza Hut", "KFC", "Domino's",
        "Burger King"], 2),
}


def _word_regex(terms):
    parts = sorted({re.escape(t).replace("'", "['’]") for t in terms}, key=len, reverse=True)
    return re.compile(r"\b(?:" + "|".join(parts) + r")s?\b", re.I)


LEXICON_RES = {cat: (_word_regex(terms), tier) for cat, (terms, tier) in LEXICONS.items()}

HONORIFIC_RE = re.compile(r"\b(?:Mr|Mrs|Ms|Mx|Dr|Prof|Miss)\b\.?(?=\s+[A-Z])")  # case-sensitive

_NUM = r"(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"      # '2,500' / '10' / '12.50'; never a trailing comma
CURRENCY_RES = [
    re.compile(r"\$\s?(?:" + _NUM + r"|\.\d+)(?:\s?(?:thousand|million|billion))?", re.I),
    re.compile(r"(?<![\w$.])" + _NUM + r"\s?(?:dollars?|cents?|bucks)\b", re.I),
    re.compile(r"(?<![\w$.])" + _NUM + r"\s?¢"),
]
CURRENCY_WORD_RE = re.compile(r"\b(?:dollars?|cents?|bucks?|nickels?|dimes?|pennies|penny)\b", re.I)

# US-customary units (a localization target when the Philippine version is metric).
UNIT_SAFE_RE = re.compile(
    r"\b(?:miles?|mph|gallons?|ounces?|inch(?:es)?|yards?|fahrenheit|acres?|"
    r"square feet|square miles)\b", re.I)
UNIT_DEG_RE = re.compile(r"°\s?F\b")
_NUMWORDS = ("a|an|one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|"
             "twenty|thirty|forty|fifty|hundred|half")
# Ambiguous words ("foot", "pound", "pint"...) only count when attached to a quantity.
UNIT_NEEDS_NUM_RE = re.compile(
    r"(?:(?<![\w$.])\d[\d,./]*[\s-]?|\b(?:" + _NUMWORDS + r")[\s-]+)"
    r"(feet|foot|ft|pounds?|lbs?|oz|pints?|quarts?)\b", re.I)

SPACY_MAP = {"PERSON": "Proper Name", "GPE": "Place", "LOC": "Place", "FAC": "Place",
             "ORG": "Organization", "EVENT": "Event", "NORP": "Nationality/Group",
             "PRODUCT": "Product", "WORK_OF_ART": "Work of Art"}
HF_MAP = {"PER": "Proper Name", "LOC": "Place", "ORG": "Organization"}  # MISC deliberately dropped

OBJECT_LEXNAMES = {"noun.food": "Food", "noun.artifact": "Object",
                   "noun.animal": "Animal", "noun.plant": "Plant"}
# If WordNet's MOST COMMON sense of a food-model hit is one of these, the word is
# almost certainly not being used as food ('birds', 'starfish', 'shoes').
FOOD_DEMOTE_FIRST = {"noun.animal", "noun.artifact", "noun.attribute", "noun.body",
                     "noun.location"}
ABSTRACT_FIRST = {"noun.act", "noun.attribute", "noun.cognition", "noun.event", "noun.state", "noun.time",
                  "noun.quantity", "noun.phenomenon", "noun.process", "noun.relation", "noun.feeling", "noun.motive"}
BAD_ENTS = {"PERSON", "DATE", "TIME", "MONEY", "CARDINAL", "QUANTITY", "PERCENT", "ORDINAL"}
QUANTIFIERS = {"how", "what", "which", "many", "much", "more", "most", "less", "least",
               "twice", "enough", "several", "few", "remaining", "only", "final", "other",
               "another", "additional", "extra", "entire", "whole", "same", "each",
               "every", "any", "no", "all", "both", "new", "old"}

PRIORITY = {"Currency": 0, "Honorific": 1, "Proper Name": 2, "Event": 3, "Unit": 4,
            "Brand": 5, "Place": 5, "Organization": 5, "Product": 5,
            "Nationality/Group": 5, "Work of Art": 5, "Food": 6, "Food?": 6,
            "Culture": 6, "Object": 7, "Animal": 7, "Plant": 7}


# ============================================================================
# DATA
# ============================================================================
@dataclass
class Cand:
    start: int
    end: int
    text: str
    category: str
    sources: set
    tier: int                       # 1 = high confidence, 2 = review, 3 = low-value
    score: Optional[float] = None   # best model confidence, if any
    lemma: Optional[str] = None     # grouping key for plurals ('ships' -> 'ship')
    pool: Optional[str] = None      # suggested replacement pool (objects/animals/plants/food from WordNet)

    @property
    def span(self):
        return (self.start, self.end)


_INST_RE = re.compile(r"^\s*(?:<s>)?\s*\[INST\]\s*(.*?)\s*\[/INST\]", re.S)
_MARKER_RE = re.compile(r"</?s>|\[/?INST\]")


def clean_problem(raw):
    """'<s>[INST]question [/INST] solution #### n </s>' -> 'question'.
    Plain question text passes through unchanged."""
    m = _INST_RE.match(raw)
    if m:
        return m.group(1).strip()
    return _MARKER_RE.sub(" ", raw).strip()


def load_problems(path=GSM8K_CSV_PATH, column=None):
    if not os.path.exists(path):
        raise FileNotFoundError(f"Could not find '{path}'.")
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            raise ValueError(f"'{path}' has no header row.")
        col = column
        if col is None:
            lower = {n.lower(): n for n in reader.fieldnames}
            col = next((lower[c] for c in QUESTION_COLUMN_CANDIDATES if c in lower),
                       reader.fieldnames[0])
        raw = [(r.get(col) or "").strip() for r in reader]
    raw = [p for p in raw if p]
    n_markup = sum("[INST]" in p for p in raw)
    problems = [p for p in (clean_problem(p) for p in raw) if p]
    if not problems:
        raise ValueError(f"No problem text found in '{path}' (column '{col}').")
    print(f"Loaded {len(problems)} problems from '{path}' (column '{col}').")
    if n_markup:
        print(f"Stripped the [INST] template from {n_markup} rows; using question text only.")
    return problems


EXTRA_NAMES = {"Weng", "Juan", "Jose", "Pedro", "Andres", "Emilio", "Ramon", "Antonio", "Manuel", "Carlos", "Miguel", "Rafael",
               "Angelo", "Paolo", "Jericho", "Renz", "Kenneth", "Noel", "Rodel", "Rey", "Ronnie", "Arnel", "Dante", "Nestor",
               "Maria", "Ana", "Rosa", "Luz", "Carmen", "Josefa", "Teresa", "Elena", "Lourdes", "Cristina", "Jasmine", "Jessa",
               "Maricel", "Marites", "Liza", "Rowena", "Nena", "Perla", "Corazon", "Imelda", "Divina", "Estrella", "Jocelyn",
               "Maribel", "Rhea", "Yolanda", "Mei", "Lin", "Wei", "Li", "Chen", "Yuki", "Hiro", "Kenji", "Priya", "Arjun", "Aisha"}


def load_term_list(path):
    if not os.path.exists(path):
        return set()
    with open(path, encoding="utf-8") as f:
        return {ln.strip().lower() for ln in f if ln.strip() and not ln.startswith("#")}


# ============================================================================
# SPAN HELPERS
# ============================================================================
def expand_to_word(text, s, e):
    """Snap a span to full word boundaries (fixes '##' / mid-word fragments)."""
    while s > 0 and text[s - 1].isalnum():
        s -= 1
    while e < len(text) and text[e].isalnum():
        e += 1
    return s, e


def cluster(spans):
    """[(start, end, source, score)] -> [[start, end, {sources}, best_score]],
    merging overlapping/touching spans."""
    out = []
    for s, e, src, sc in sorted(spans, key=lambda t: (t[0], t[1], t[2])):
        if out and s <= out[-1][1]:
            out[-1][1] = max(out[-1][1], e)
            out[-1][2].add(src)
            if sc is not None and (out[-1][3] is None or sc > out[-1][3]):
                out[-1][3] = sc
        else:
            out.append([s, e, {src}, sc])
    return out


_HONOR_PREFIX = re.compile(r"^(?:Mr|Mrs|Ms|Mx|Dr|Prof|Miss)\b\.?\s+")
_POSSESSIVE = re.compile(r"['’]s$")


def trim_person(text, s, e):
    """Drop a leading honorific ('Dr. Smith' -> 'Smith') and trailing possessive."""
    m = _HONOR_PREFIX.match(text[s:e])
    if m:
        s += m.end()
    m = _POSSESSIVE.search(text[s:e])
    if m:
        e -= len(m.group())
    return s, e


def resolve(cands):
    """Overlap resolution: highest-priority category wins, then longest span.
    Identical spans+category from different tools are merged (sources union)."""
    cands = sorted(cands, key=lambda c: (PRIORITY.get(c.category, 8), -(c.end - c.start), c.start))
    kept = []
    for c in cands:
        clash = [k for k in kept if c.start < k.end and k.start < c.end]
        if not clash:
            kept.append(c)
        elif len(clash) == 1 and clash[0].span == c.span and clash[0].category == c.category:
            clash[0].sources |= c.sources
            clash[0].tier = min(clash[0].tier, c.tier)
            if c.score is not None and (clash[0].score is None or c.score > clash[0].score):
                clash[0].score = c.score
    return sorted(kept, key=lambda c: c.start)


# ============================================================================
# WORDNET HELPERS
# ============================================================================
@lru_cache(maxsize=None)
def wn_category(lemma):
    """Generic-object classification from the first WN_SENSES senses."""
    if wn is None:
        return None
    syns = wn.synsets(lemma, pos="n")
    if syns and syns[0].lexname() in ABSTRACT_FIRST:
        return None
    for syn in syns[:WN_SENSES]:
        cat = OBJECT_LEXNAMES.get(syn.lexname())
        if cat:
            return cat
    return None


def is_object_noun(tok):
    if tok.is_stop or tok.ent_type_ in BAD_ENTS or not tok.text.isalpha():
        return False
    return not ({tok.text.lower(), tok.lemma_.lower()} & EXCLUDED_NOUNS)


@lru_cache(maxsize=None)
def food_verdict(span_text):
    """Decide what to do with a food-model hit. Returns (category, tier) or None (drop).

    - generic words ('food', 'meal')            -> dropped
    - word unknown to WordNet                   -> Food, Tier 2 (could be a local dish)
    - most common sense is food                 -> Food, Tier 2
    - most common sense is animal/artifact/...  -> Food?, Tier 3 ('birds', 'shoes')
    - otherwise Tier 2 only if a food sense appears in the top 4, else Food?, Tier 3
      (keeps 'brownies', whose food sense is not first; drops flowers/colors like 'violet')
    """
    words = re.findall(r"[A-Za-z'-]+", span_text.lower())
    if not words:
        return None
    head = words[-1]
    if head in GENERIC_FOOD:
        return None
    if wn is None:
        return ("Food", 2)
    syns = wn.synsets("_".join(words), pos="n") if len(words) > 1 else []
    if not syns:
        syns = wn.synsets(wn.morphy(head, "n") or head, pos="n")
    if not syns:
        return ("Food", 2)
    first = syns[0].lexname()
    if first == "noun.food":
        return ("Food", 2)
    if first in FOOD_DEMOTE_FIRST:
        return ("Food?", 3)
    if any(s.lexname() == "noun.food" for s in syns[:4]):
        return ("Food", 2)
    return ("Food?", 3)


def weak_single_word(doc, s, e, srcs):
    """One-token ORG/Place with a single vote whose POS is not PROPN ('Next', 'Then', 'Today'): a common word, not an entity."""
    if len(srcs) > 1 or not hasattr(doc, "char_span"):
        return False
    span = doc.char_span(s, e, alignment_mode="expand")
    return span is not None and len(span) == 1 and span[0].pos_ not in ("PROPN", "X", "")


def looks_like_person(doc, s, e, word, names):
    """Single-word ORG/Place guess that is really a name: a known first name,
    or the grammatical subject of a verb ('Jerry sold ...')."""
    if " " in word.strip() or LEXICON_RES["Brand"][0].fullmatch(word):
        return False, False
    if word in names:
        return True, True            # (is_person, in_name_list)
    char_span = doc.char_span(s, e, alignment_mode="expand") if hasattr(doc, "char_span") else None
    if char_span is not None and len(char_span) == 1:
        tok = char_span[0]
        return (tok.dep_ in ("nsubj", "nsubjpass") and tok.head.pos_ == "VERB"), False
    return False, False


# ============================================================================
# SUPPORT VOTES FOR OBJECTS (independent evidence beyond WordNet)
# ============================================================================
# WordNet hypernym stem -> pool, most specific first (first match wins)
HYPERNYM_POOLS = [("domestic_animal", "Pets"), ("bird", "Birds"), ("insect", "Insects & small creatures"),
                  ("fish", "Aquatic animals"), ("shellfish", "Aquatic animals"), ("mammal", "Wild animals"),
                  ("clothing", "Clothing & accessories"), ("garment", "Clothing & accessories"),
                  ("footwear", "Clothing & accessories"), ("container", "Containers & packaging"),
                  ("furniture", "Household & furniture"), ("kitchenware", "Household & furniture"),
                  ("tableware", "Household & furniture"), ("vehicle", "Vehicles & transport"),
                  ("plaything", "Toys, games & sports"), ("sports_equipment", "Toys, games & sports"),
                  ("building", "Places & buildings"), ("tree", "Plants & nature"), ("flower", "Plants & nature"),
                  ("fruit", "Fruit"), ("vegetable", "Vegetable"), ("beverage", "Beverages"),
                  ("baked_goods", "Baked goods"), ("dairy_product", "Eggs & dairy"), ("meat", "Meat & poultry"),
                  ("candy", "Candy & sweets")]
QUANT_CHILD = {"many", "few", "several", "couple", "dozen", "each", "every", "more"}
MEASURE_NOUNS = {"box", "pack", "bag", "pair", "dozen", "bunch", "jar", "can", "bottle", "carton", "slice", "piece", "cup",
                 "pound", "gallon", "tray", "crate", "basket", "bowl", "plate", "stack", "set", "group", "row", "case"}


def _forms(word):
    w = word.lower(); out = [w]
    if w.endswith("ies"): out.append(w[:-3] + "y")
    if w.endswith("es"): out.append(w[:-2])
    if w.endswith("s"): out.append(w[:-1])
    return out


@lru_cache(maxsize=None)
def hypernym_pool(lemma):
    """Suggested pool from the WordNet hypernym path ('shirt' -> garment -> clothing). Not an independent vote."""
    if wn is None:
        return None
    try:
        for syn in wn.synsets(lemma, pos="n")[:WN_SENSES]:
            stems = {syn.name().split(".")[0]} | {h.name().split(".")[0] for h in syn.closure(lambda x: x.hypernyms())}
            for stem, pool in HYPERNYM_POOLS:
                if stem in stems:
                    return pool
    except Exception:
        return None
    return None


def is_counted(tok):
    """The problem itself counts/measures this noun: '7 starfish', 'how many gnomes', 'a box of cookies'."""
    kids = list(getattr(tok, "children", ()))
    if any(c.dep_ == "nummod" or c.text.lower() in QUANT_CHILD for c in kids):
        return True
    head = getattr(tok, "head", None)
    return bool(tok.dep_ == "pobj" and head is not None and head.text.lower() == "of"
                and getattr(head, "head", None) is not None and head.head.lemma_.lower() in MEASURE_NOUNS)


def is_concrete(lemma):
    return bool(CONCRETENESS) and CONCRETENESS.get(lemma.split()[-1], 0) >= CONCRETE_MIN


def load_concreteness(path):
    """Brysbaert et al. (2014) concreteness ratings: columns 'Word' and 'Conc.M' (.txt/.tsv/.csv/.xlsx)."""
    rows = []
    if path.lower().endswith(".xlsx"):
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, encoding="utf-8-sig", newline="") as f:
            first = f.readline(); f.seek(0)
            rows = list(csv.reader(f, delimiter="\t" if "\t" in first else ","))
    head = [str(h).strip().lower() for h in rows[0]]
    wi, ci = head.index("word"), head.index("conc.m")
    for r in rows[1:]:
        try:
            CONCRETENESS[str(r[wi]).strip().lower()] = float(r[ci])
        except (ValueError, TypeError, IndexError):
            continue
    print(f"Loaded {len(CONCRETENESS)} concreteness ratings from {path}")


# ============================================================================
# CANDIDATE COLLECTION
# ============================================================================
def collect(text, doc, hf_ents, food_ents, names=frozenset()):
    cands = []

    def add(s, e, cat, src, tier, score=None):
        if s < e:
            cands.append(Cand(s, e, text[s:e], cat, {src}, tier, score))

    # ---- rule-based, high precision ----
    for m in HONORIFIC_RE.finditer(text):
        add(m.start(), m.end(), "Honorific", "regex", 1)
    for rx in CURRENCY_RES:
        for m in rx.finditer(text):
            add(m.start(), m.end(), "Currency", "regex", 1)
    for m in CURRENCY_WORD_RE.finditer(text):
        add(m.start(), m.end(), "Currency", "regex-word", 2)
    for rx in (UNIT_SAFE_RE, UNIT_DEG_RE):
        for m in rx.finditer(text):
            add(m.start(), m.end(), "Unit", "regex", 1)
    for m in UNIT_NEEDS_NUM_RE.finditer(text):
        add(m.start(1), m.end(1), "Unit", "regex", 1)
    for cat, (rx, tier) in LEXICON_RES.items():
        for m in rx.finditer(text):
            add(m.start(), m.end(), cat, "lexicon", tier)

    # ---- NER with voting (spaCy + transformer + gazetteer) ----
    spans = defaultdict(list)
    for ent in doc.ents:
        cat = SPACY_MAP.get(ent.label_)
        if cat:
            spans[cat].append((ent.start_char, ent.end_char, "spacy", None))
    for ent in hf_ents:
        cat = HF_MAP.get(ent.get("entity_group"))
        if cat and ent["score"] >= NER_MIN_SCORE:
            s, e = expand_to_word(text, ent["start"], ent["end"])
            spans[cat].append((s, e, "hf", float(ent["score"])))
    if names:
        for tok in doc:
            if tok.pos_ == "PROPN" and tok.text.istitle() and tok.text in names:
                spans["Proper Name"].append((tok.idx, tok.idx + len(tok.text), "gazetteer", None))

    for cat, items in spans.items():
        for s, e, srcs, sc in cluster(items):
            if cat == "Proper Name":
                s, e = trim_person(text, s, e)
                if s >= e:
                    continue
            word = text[s:e]
            has_model = bool(srcs & {"spacy", "hf"})
            # "May", "April", "Monday"... are names only if a real NER model says so.
            if word.lower() in (MONTHS | DAYS) and not (cat == "Proper Name" and has_model):
                continue
            srcs = set(srcs)
            if cat in ("Organization", "Place"):
                person, in_names = looks_like_person(doc, s, e, word, names)
                if person:
                    cat = "Proper Name"
                    srcs.add("gazetteer" if in_names else "reassigned")
            if cat in ("Organization", "Place", "Nationality/Group", "Product", "Work of Art") and weak_single_word(doc, s, e, srcs):
                continue
            votes = len(srcs - {"reassigned"})
            cands.append(Cand(s, e, word, cat, srcs, 1 if votes >= 2 else 2, sc))

    # ---- food model, gated by WordNet ----
    food_spans = []
    for ent in food_ents:
        if ent["score"] >= FOOD_MIN_SCORE:
            s, e = expand_to_word(text, ent["start"], ent["end"])
            food_spans.append((s, e, "foodbert", float(ent["score"])))
    for s, e, _, sc in cluster(food_spans):
        verdict = food_verdict(text[s:e].strip())
        if verdict:
            cat, tier = verdict
            cands.append(Cand(s, e, text[s:e], cat, {"foodbert"}, tier, sc))

    # ---- generic objects via WordNet (low tier) ----
    for tok in doc:
        if tok.pos_ != "NOUN" or not is_object_noun(tok):
            continue
        cat = wn_category(tok.lemma_.lower())
        if not cat:
            continue
        start = tok.i
        # extend left over noun-noun compounds only ("care package", "boarding school")
        while (start > 0 and doc[start - 1].dep_ == "compound"
               and doc[start - 1].head.i >= start
               and doc[start - 1].text.isalpha()
               and doc[start - 1].pos_ in ("NOUN", "VERB", "ADJ")
               and doc[start - 1].text.lower() not in QUANTIFIERS):
            start -= 1
        add(doc[start].idx, tok.idx + len(tok.text), cat, "wordnet", 2 if cat == "Food" else 3)
        c = cands[-1]
        c.lemma = " ".join([t.text.lower() for t in doc[start:tok.i]] + [tok.lemma_.lower()])
        if is_counted(tok):
            c.sources.add("counted")
        if is_concrete(tok.lemma_):
            c.sources.add("concrete")
        c.pool = hypernym_pool(tok.lemma_.lower())
        if cat != "Food" and len(c.sources) >= 2:
            c.tier = 2                       # WordNet + at least one independent vote

    if REJECT_TERMS:
        cands = [c for c in cands if c.text.lower() not in REJECT_TERMS]
    return resolve(cands)


# ============================================================================
# MODEL LOADING
# ============================================================================
def load_spacy():
    import spacy
    for name in (SPACY_MODEL, "en_core_web_sm"):
        try:
            nlp = spacy.load(name)
            print(f"spaCy model: {name}")
            return nlp
        except Exception as exc:
            print(f"[warn] could not load spaCy model {name}: {exc}")
    raise SystemExit("No spaCy model available. Run: python -m spacy download en_core_web_sm")


def load_hf(model_name):
    if hf_pipeline is None:
        print(f"[warn] transformers not installed; skipping {model_name}")
        return None
    try:
        import torch
        device = 0 if torch.cuda.is_available() else -1
    except ImportError:
        device = -1
    return hf_pipeline("token-classification", model=model_name,
                       aggregation_strategy="simple", device=device)


def run_hf(pipe, texts):
    if pipe is None:
        return [[] for _ in texts]
    try:
        return pipe(texts, batch_size=HF_BATCH)
    except Exception as exc:
        print(f"[warn] HF pipeline failed ({exc}); continuing without it")
        return [[] for _ in texts]


def load_nltk_resources():
    names = set()
    if nltk is None:
        print("[warn] nltk not installed: no NLTK name list, no WordNet filters")
        return EXTRA_NAMES | {n.title() for n in load_term_list(EXTRA_NAMES_FILE)}
    for pkg in ("wordnet", "names"):
        try:
            wn.synsets("dog") if pkg == "wordnet" else nltk_names.words()
        except LookupError:
            nltk.download(pkg, quiet=True)
    try:
        names = set(nltk_names.words())
    except LookupError:
        print("[warn] NLTK 'names' corpus unavailable")
    return names | EXTRA_NAMES | {n.title() for n in load_term_list(EXTRA_NAMES_FILE)}


# ============================================================================
# REPORTING
# ============================================================================
HEADERS = ["Problem #", "Candidate", "Context", "Category", "Tier", "Sources", "Score",
           "Count", "Suggested Pool", "Decision", "Final Category", "Localized Replacement"]
TIER_FILL = {1: "E2EFDA", 2: "FFF2CC", 3: "EDEDED"}


def context_snippet(text, c, width=35):
    s, e = max(0, c.start - width), min(len(text), c.end + width)
    return ("…" if s else "") + text[s:c.start] + "[" + text[c.start:c.end] + "]" + \
           text[c.end:e] + ("…" if e < len(text) else "")


def build_rows(problems, results):
    rows = []
    for i, (text, cands) in enumerate(zip(problems, results), 1):
        groups = {}
        for c in cands:  # one row per unique (term, category) per problem
            key = (c.lemma or c.text.lower(), c.category)
            if key in groups:
                g = groups[key]
                g[1] += 1
                g[0].sources |= c.sources
                g[0].tier = min(g[0].tier, c.tier)
                if c.score is not None and (g[0].score is None or c.score > g[0].score):
                    g[0].score = c.score
                g[0].pool = g[0].pool or c.pool
            else:
                groups[key] = [c, 1]
        for c, count in sorted(groups.values(), key=lambda g: (g[0].tier, g[0].start)):
            rows.append((i, c.text, context_snippet(text, c), c.category, c.tier,
                         ", ".join(sorted(c.sources)),
                         None if c.score is None else round(c.score, 2), count, c.pool or ""))
    return rows


def write_report(problems, results, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    ws.append(HEADERS)
    rows = build_rows(problems, results)
    for r in rows:
        ws.append(list(r) + ["", "", ""])   # Decision, Final Category, Localized Replacement

    head_fill = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    tier_ix = HEADERS.index("Tier")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        t = row[tier_ix].value
        if t in TIER_FILL:
            row[tier_ix].fill = PatternFill("solid", start_color=TIER_FILL[t], end_color=TIER_FILL[t])
    widths = (10, 22, 60, 18, 6, 26, 7, 7, 24, 11, 22, 24)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if ws.max_row >= 2:
        col = get_column_letter(HEADERS.index("Decision") + 1)
        dv = DataValidation(type="list", formula1='"Keep,Reject,Unsure"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")

    ws2 = wb.create_sheet("Per Problem (Tier 1-2)")
    ws2.append(["Problem #", "Problem Text", "Candidates (Tier 1-2)"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    by_problem = defaultdict(lambda: defaultdict(list))
    for i, term, _, cat, tier, *_ in rows:
        if tier <= 2:
            by_problem[i][cat].append(term)
    for i, text in enumerate(problems, 1):
        lines = [f"{cat}: {', '.join(terms)}" for cat, terms in sorted(by_problem[i].items())]
        ws2.append([i, text, "\n".join(lines) or "(none)"])
        r = ws2.max_row
        for c in range(1, 4):
            ws2.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        ws2.row_dimensions[r].height = max(30, 15 * max(1, len(lines)))
    ws2.column_dimensions["A"].width = 10
    ws2.column_dimensions["B"].width = 60
    ws2.column_dimensions["C"].width = 60
    ws2.freeze_panes = "A2"

    wb.save(path)
    print(f"Saved: {path}")



def review_generated_xlsx(path):
    """Promote recurring exact Candidate/Category matches after the workbook is saved.

    A term is eligible when it has at least three rows at Tier 1 or Tier 2,
    and also appears at Tier 3, all with exactly matching Candidate and
    Category cell values. Eligible Tier 3 rows are raised to Tier 2.
    """
    wb = load_workbook(path)
    if "Candidates" not in wb.sheetnames:
        print("[warn] Could not review workbook: missing 'Candidates' sheet.")
        wb.close()
        return

    ws = wb["Candidates"]
    headers = {cell.value: cell.column for cell in ws[1]}
    required = ("Candidate", "Category", "Tier", "Problem #")
    if any(name not in headers for name in required):
        print("[warn] Could not review workbook: required columns are missing.")
        wb.close()
        return

    candidate_col = headers["Candidate"]
    category_col = headers["Category"]
    tier_col = headers["Tier"]
    problem_col = headers["Problem #"]

    # Count rows at Tier 1 or 2 by exact cell values; do not lowercase,
    # strip, normalize punctuation, or otherwise alter Candidate/Category.
    high_tier_counts = defaultdict(int)
    has_tier3 = set()
    for row_num in range(2, ws.max_row + 1):
        candidate = ws.cell(row_num, candidate_col).value
        category = ws.cell(row_num, category_col).value
        tier = ws.cell(row_num, tier_col).value
        if candidate is None or category is None:
            continue
        key = (candidate, category)
        if isinstance(tier, (int, float)) and tier <= 2:
            high_tier_counts[key] += 1
        elif tier == 3:
            has_tier3.add(key)

    eligible = {
        key for key, count in high_tier_counts.items()
        if count >= 3 and key in has_tier3
    }

    promoted = 0
    tier_fill = PatternFill("solid", start_color=TIER_FILL[2], end_color=TIER_FILL[2])
    for row_num in range(2, ws.max_row + 1):
        candidate = ws.cell(row_num, candidate_col).value
        category = ws.cell(row_num, category_col).value
        tier_cell = ws.cell(row_num, tier_col)
        if (candidate, category) in eligible and tier_cell.value == 3:
            tier_cell.value = 2
            tier_cell.fill = tier_fill
            promoted += 1

    # Refresh the derived "Per Problem (Tier 1-2)" sheet so promoted rows
    # are included there too, while preserving each problem's original text.
    if "Per Problem (Tier 1-2)" in wb.sheetnames:
        summary = wb["Per Problem (Tier 1-2)"]
        summary_headers = {cell.value: cell.column for cell in summary[1]}
        if all(name in summary_headers for name in ("Problem #", "Problem Text", "Candidates (Tier 1-2)")):
            problem_texts = {}
            for row_num in range(2, summary.max_row + 1):
                problem_num = summary.cell(row_num, summary_headers["Problem #"]).value
                problem_text = summary.cell(row_num, summary_headers["Problem Text"]).value
                problem_texts[problem_num] = problem_text

            by_problem = defaultdict(lambda: defaultdict(list))
            for row_num in range(2, ws.max_row + 1):
                problem_num = ws.cell(row_num, problem_col).value
                candidate = ws.cell(row_num, candidate_col).value
                category = ws.cell(row_num, category_col).value
                tier = ws.cell(row_num, tier_col).value
                if problem_num is not None and candidate is not None and isinstance(tier, (int, float)) and tier <= 2:
                    by_problem[problem_num][category].append(candidate)

            text_col = summary_headers["Problem Text"]
            candidates_col = summary_headers["Candidates (Tier 1-2)"]
            for row_num in range(2, summary.max_row + 1):
                problem_num = summary.cell(row_num, summary_headers["Problem #"]).value
                summary.cell(row_num, text_col).value = problem_texts.get(problem_num, "")
                lines = [
                    f"{cat}: {', '.join(terms)}"
                    for cat, terms in sorted(by_problem[problem_num].items(), key=lambda item: str(item[0]))
                ]
                summary.cell(row_num, candidates_col).value = "\\n".join(lines) or "(none)"
                summary.cell(row_num, text_col).alignment = Alignment(wrap_text=True, vertical="top")
                summary.cell(row_num, candidates_col).alignment = Alignment(wrap_text=True, vertical="top")
                summary.row_dimensions[row_num].height = max(30, 15 * max(1, len(lines)))

    wb.save(path)
    wb.close()
    print(
        f"Post-generation review: promoted {promoted} Tier 3 row(s) to Tier 2 "
        f"across {len(eligible)} exact Candidate/Category match(es)."
    )

def _decision(value):
    d = str(value or "").strip().lower()
    return d if d in ("keep", "reject") else None


def _read_reviewed(path):
    ws = load_workbook(path, read_only=True)["Candidates"]
    rows = ws.iter_rows(values_only=True)
    ix = {h: i for i, h in enumerate(next(rows))}
    return ix, rows


def precision_report(path):
    """Precision (Keep / (Keep + Reject)) per category, tier and source tool,
    from a reviewed workbook. Recall needs a hand-labeled gold set."""
    ix, rows = _read_reviewed(path)
    tables = {"category": defaultdict(lambda: [0, 0]),
              "tier": defaultdict(lambda: [0, 0]),
              "source": defaultdict(lambda: [0, 0])}
    read = counted = 0
    for r in rows:
        if r[ix["Candidate"]] is None:
            continue
        read += 1
        d = _decision(r[ix["Decision"]])
        if d is None:
            continue
        counted += 1
        k = 0 if d == "keep" else 1
        tables["category"][r[ix["Category"]]][k] += 1
        tables["tier"][r[ix["Tier"]]][k] += 1
        for s in str(r[ix["Sources"]]).split(", "):
            tables["source"][s][k] += 1
    print(f"Rows read: {read} | with a Keep/Reject decision: {counted} | undecided: {read - counted}")
    if not counted:
        print("No Keep/Reject decisions found. Fill the Decision column (use the dropdown), "
              "save the file, close Excel, and run again.")
        return
    for name, table in tables.items():
        print(f"\nPrecision by {name}")
        print(f"{'':24}{'keep':>6}{'reject':>8}{'precision':>11}")
        for key, (keep, rej) in sorted(table.items(), key=lambda kv: str(kv[0])):
            print(f"{str(key):24}{keep:>6}{rej:>8}{keep / (keep + rej):>11.0%}")


def learn_rejects(path, min_rejects=2):
    """Add terms you rejected >= min_rejects times and NEVER kept to reject_terms.txt."""
    ix, rows = _read_reviewed(path)
    tally = defaultdict(lambda: [0, 0])
    for r in rows:
        d = _decision(r[ix["Decision"]])
        if d is None or r[ix["Candidate"]] is None:
            continue
        tally[str(r[ix["Candidate"]]).strip().lower()][0 if d == "keep" else 1] += 1
    new = {t for t, (k, rej) in tally.items() if k == 0 and rej >= min_rejects}
    merged = sorted(load_term_list(REJECT_FILE) | new)
    with open(REJECT_FILE, "w", encoding="utf-8") as f:
        f.write("# one term per line (lowercase); hidden from future runs\n")
        f.write("\n".join(merged) + "\n")
    print(f"Added {len(new)} new terms; {len(merged)} total in {REJECT_FILE}")


# ============================================================================
# MAIN
# ============================================================================
def extract(problems):
    REJECT_TERMS.update(load_term_list(REJECT_FILE))
    if REJECT_TERMS:
        print(f"Loaded {len(REJECT_TERMS)} reject terms from {REJECT_FILE}")
    nlp = load_spacy()
    names = load_nltk_resources()
    hf_ner = load_hf(HF_NER_MODEL)
    food_ner = load_hf(FOOD_MODEL)

    print("Running spaCy...")
    docs = list(nlp.pipe(problems, batch_size=64))
    print("Running transformer NER...")
    hf_out = run_hf(hf_ner, problems)
    print("Running food NER...")
    food_out = run_hf(food_ner, problems)
    return [collect(t, d, h, f, names) for t, d, h, f in zip(problems, docs, hf_out, food_out)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default=GSM8K_CSV_PATH)
    ap.add_argument("--column")
    ap.add_argument("--out", default="gsmph_candidates.xlsx")
    ap.add_argument("--limit", type=int, help="only process the first N problems")
    ap.add_argument("--concreteness", metavar="FILE", help="Brysbaert concreteness ratings (.txt/.csv/.xlsx): extra object vote")
    ap.add_argument("--precision-report", metavar="REVIEWED_XLSX")
    ap.add_argument("--learn-rejects", metavar="REVIEWED_XLSX")
    args = ap.parse_args()

    if args.precision_report:
        precision_report(args.precision_report)
        return
    if args.learn_rejects:
        learn_rejects(args.learn_rejects)
        return
    if args.concreteness:
        load_concreteness(args.concreteness)
    problems = load_problems(args.csv, args.column)
    if args.limit:
        problems = problems[:args.limit]
    write_report(problems, extract(problems), args.out)
    review_generated_xlsx(args.out)


if __name__ == "__main__":
    main()
