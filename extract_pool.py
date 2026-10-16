#!/usr/bin/env python3
"""
extract_pool.py - build replacement pools of Filipino terms from open, citable sources.

Sources (all free APIs, no scraping of HTML pages):
  * Wikipedia / Wiktionary category trees  (MediaWiki API)    -> candidate titles + short descriptions
  * Wikidata                               (wbgetentities)    -> Tagalog label, aliases, instance-of, country of origin
  * Wikimedia Pageviews                    (REST API)         -> how recognizable each term is (12-month views)

Output: one Excel workbook (Pool / Sources / Run Log / README) and, optionally, a CSV.

Usage
  pip install requests openpyxl
  python extract_pool.py list-pools
  python extract_pool.py discover --contact you@example.com        # find real category names first
  python extract_pool.py run --contact you@example.com --out pool.xlsx
  python extract_pool.py run --contact you@example.com --pools "Baked goods,Fruit" --no-views

Generic nouns (store, school, house, box, milk, dog ...) come from a Tagalog dictionary, not Wikipedia:
  python extract_pool.py fetch-dict --contact you@example.com                 # kaikki.org Tagalog Wiktionary extract
  python extract_pool.py discover-dict kaikki_tagalog.jsonl                   # see real category/topic names, tune DICT_RULES
  python extract_pool.py targets --targets GSMPH_Cleaned_Terms.xlsx           # how many terms each pool needs (+10%)
  python extract_pool.py run --contact you@example.com --dict kaikki_tagalog.jsonl --names-file names.csv --out pool.xlsx
The workbook gets Coverage (pool size vs term-bank need, with formulas), Generic Terms and Rejected (auto) sheets.

The category names in POOLS below are SEEDS written without network access. Run `discover`
first, then edit POOLS to match the categories that actually exist. The Run Log sheet lists
every category that returned 0 members so nothing fails silently.

Wikimedia asks API users to send a descriptive User-Agent with contact info: use --contact.
Text from Wikipedia/Wiktionary is CC BY-SA 4.0 and Wikidata is CC0: keep the Source/URL
columns if you redistribute the pool.
"""
import argparse, csv, datetime as dt, gzip, hashlib, json, math, re, sys, time, unicodedata
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import quote

import requests

# --------------------------------------------------------------------------------------
# CONFIG: pools keyed by the same names as the term-bank subcategories
# S(category, depth, host)  depth = how many levels of sub-categories to follow
# must = regex that title/description/extract must match (precision filter)
# exclude = regex that rejects a page (None disables the default exclusion)
# --------------------------------------------------------------------------------------
EN = "en.wikipedia.org"
def S(category, depth=1, host=EN): return dict(host=host, category=category, depth=depth)

DEFAULT_EXCLUDE = (r"\b(politician|footballer|basketball player|actor|actress|singer|rapper|album|"
                   r"film|television series|novel|song|businessman|businesswoman)\b")

POOLS = {
 "Baked goods": dict(sources=[S("Philippine desserts", 1), S("Philippine breads", 1)],
                     must=r"cake|bread|pastry|bun|biscuit|cookie|kakanin|rice cake|sponge|pudding|glutinous|steamed"),
 "Frozen treats": dict(sources=[S("Philippine desserts", 1)], must=r"ice|frozen|sorbet|shaved|halo|gelato"),
 "Candy & sweets": dict(sources=[S("Philippine desserts", 1), S("Philippine confectionery", 1)],
                        must=r"candy|sweet|confection|caramel|milk|polvor|sugar"),
 "Meals & fast food": dict(sources=[S("Philippine cuisine", 2)],
                           must=r"dish|soup|noodle|stew|rice|dumpling|roll|sausage|fried|grilled|sandwich"),
 "Snacks": dict(sources=[S("Philippine snack foods", 1), S("Philippine street food", 1)], must=None),
 "Beverages": dict(sources=[S("Philippine drinks", 1)], must=None),
 "Fruit": dict(sources=[S("Fruits originating in Asia", 0), S("Philippine fruits", 1)],
               must=r"fruit|tree|berry"),
 "Vegetable": dict(sources=[S("Philippine vegetables", 1)], must=None),
 "Seafood": dict(sources=[S("Philippine seafood dishes", 1), S("Fish of the Philippines", 1)], must=None),
 "Farm animals": dict(sources=[S("Domesticated animals of the Philippines", 1), S("Livestock in the Philippines", 1)], must=None),
 "Wild animals": dict(sources=[S("Mammals of the Philippines", 1), S("Reptiles of the Philippines", 1)], must=None),
 "Birds": dict(sources=[S("Birds of the Philippines", 1)], must=None),
 "Insects & small creatures": dict(sources=[S("Insects of the Philippines", 1)], must=None),
 "Plants & nature": dict(sources=[S("Trees of the Philippines", 1), S("Flora of the Philippines", 1)], must=None),
 "Clothing & accessories": dict(sources=[S("Philippine clothing", 1), S("Philippine traditional clothing", 1)], must=None),
 "Vehicles & transport": dict(sources=[S("Transportation in the Philippines", 1), S("Road transport in the Philippines", 1)],
                              must=r"jeepney|tricycle|bus|train|ferry|boat|bangka|rail|vehicle|taxi|bike|pedicab"),
 "Toys, games & sports": dict(sources=[S("Philippine games", 1), S("Sports in the Philippines", 1)], must=None),
 "Household & furniture": dict(sources=[S("Philippine handicrafts", 1), S("Philippine furniture", 1)], must=None),
 "Place - US/foreign city or state": dict(sources=[S("Provinces of the Philippines", 1), S("Cities of the Philippines", 1)],
                                          must=None, exclude=None),
 "Holiday/event": dict(sources=[S("Philippine festivals", 1), S("Public holidays in the Philippines", 1)], must=None, exclude=None),
 "Brand/company": dict(sources=[S("Fast-food chains of the Philippines", 1), S("Supermarkets of the Philippines", 1),
                                S("Retail companies of the Philippines", 1)], must=None, exclude=None),
 "Person name": dict(sources=[S("Filipino given names", 1), S("Tagalog given names", 1)], must=None, exclude=None, rec=False),
 # ---- pools below were missing: every non-currency Replacement Pool in the term bank now has a source ----
 # sources=[] means "no Wikipedia category fits"; these pools are filled from the Tagalog dictionary (--dict) + CURATED.
 "Places & buildings": dict(sources=[], rec=False),
 "Containers & packaging": dict(sources=[], rec=False),
 "Staples & ingredients": dict(sources=[], rec=False),
 "Eggs & dairy": dict(sources=[], rec=False),
 "Meat & poultry": dict(sources=[], rec=False),
 "Pets": dict(sources=[], rec=False),
 "Currency word": dict(sources=[], rec=False),
 "Aquatic animals": dict(sources=[S("Fish of the Philippines", 1)], must=None),
 "Place - country/region": dict(sources=[S("Regions of the Philippines", 0), S("Island groups of the Philippines", 1)], must=None, exclude=None),
 "Place/institution name": dict(sources=[S("Universities and colleges in the Philippines", 0), S("Restaurant chains in the Philippines", 1)], must=None, exclude=None),
 "Institution/government": dict(sources=[S("Government agencies of the Philippines", 1)], must=None, exclude=None),
 "Street/road name": dict(sources=[S("Streets in Metro Manila", 1), S("Roads in the Philippines", 1)], must=None, exclude=None),
}

# Currency AMOUNT pools are converted by rule (scale + realistic PH prices), they are not term pools.
RULE_BASED = {"Currency - small whole amount (<=20)", "Currency - medium whole amount (21-200)",
              "Currency - large whole amount (>200)", "Currency - decimal amount (centavos)"}

# Tagalog-dictionary matching for generic nouns (kaikki.org Wiktionary extract, noun entries only).
#   cats  = regex on the entry's Wiktionary categories + topics (lower-cased text)
#   gloss = regex on the first two English glosses
# Both are starting points: run `discover-dict` to see the real category/topic names and tune them.
DICT_RULES = {
 "Places & buildings": dict(cats=r"\b(buildings?|shops?|stores?|markets?|farms?|schools?|houses?|dwellings?)\b",
     gloss=r"\b(store|shop|market|school|house|home|farm|garden|church|town|village|restaurant|eatery|bakery|library|hospital|building|plaza|park|yard|office|beach|pond|neighborhood)\b"),
 "Containers & packaging": dict(cats=r"\b(containers?|baskets?|bottles?|vessels?)\b",
     gloss=r"\b(container|basket|bucket|pail|bottle|jar|sack|bag|box|can|tin|pot|jug|bowl|tray|packet|wrapper|crate)\b"),
 "Household & furniture": dict(cats=r"\b(furniture|kitchenware|tableware|household|bedding|utensils?)\b",
     gloss=r"\b(table|chair|bed|mat|pillow|blanket|cabinet|shelf|mirror|broom|plate|spoon|fork|cup|glass|bowl|pan|stove|fan|lamp|mosquito net)\b"),
 "Plants & nature": dict(cats=r"\b(plants?|trees?|flowers?|palms?|bamboo|grasses)\b", gloss=r"\b(tree|flower|plant|palm|bamboo|grass|shrub|vine)\b"),
 "Fruit": dict(cats=r"\bfruits?\b", gloss=r"\b(fruit|banana|mango|coconut|citrus|berry)\b"),
 "Vegetable": dict(cats=r"\bvegetables?\b", gloss=r"\b(vegetable|squash|eggplant|gourd|bean|leafy|yam|taro|cassava|sweet potato)\b"),
 "Birds": dict(cats=r"\bbirds?\b", gloss=r"\b(bird|pigeon|dove|parrot|eagle|owl|sparrow|heron|kingfisher)\b"),
 "Farm animals": dict(cats=r"\b(livestock|domestic animals?)\b", gloss=r"\b(pig|hog|cow|ox|carabao|water buffalo|goat|sheep|horse|chicken|hen|rooster|duck|goose)\b"),
 "Wild animals": dict(cats=r"\b(mammals?|reptiles?|snakes?|primates?|rodents?)\b", gloss=r"\b(monkey|deer|boar|bat|snake|lizard|crocodile|tarsier|python|civet)\b"),
 "Insects & small creatures": dict(cats=r"\b(insects?|arachnids?|spiders?|worms?)\b",
     gloss=r"\b(insect|ant|bee|fly|mosquito|cockroach|beetle|butterfly|spider|worm|cricket|grasshopper|dragonfly)\b"),
 "Aquatic animals": dict(cats=r"\b(fish|crustaceans?|molluscs?|shellfish|marine)\b",
     gloss=r"\b(fish|crab|shrimp|prawn|squid|octopus|oyster|mussel|clam|eel|shark|ray|sea turtle|whale|dolphin)\b"),
 "Clothing & accessories": dict(cats=r"\b(clothing|garments?|footwear|headwear)\b",
     gloss=r"\b(clothes|clothing|shirt|dress|skirt|pants|trousers|shoe|slipper|sandal|hat|cap|scarf|belt)\b"),
 "Vehicles & transport": dict(cats=r"\b(vehicles?|boats?|watercraft)\b", gloss=r"\b(vehicle|boat|ship|car|bus|jeep|cart|carriage|bicycle|train|plane|canoe|raft)\b"),
 "Toys, games & sports": dict(cats=r"\b(games?|toys?|sports?)\b", gloss=r"\b(game|toy|kite|marbles|sport|ball)\b"),
 "Events & activities": dict(cats=r"\b(festivals?|holidays?|celebrations?)\b",
     gloss=r"\b(festival|feast|party|celebration|wedding|birthday|baptism|holiday|fiesta|contest|concert|picnic|reunion|funeral|wake)\b"),
 "Staples & ingredients": dict(cats=r"\b(condiments?|spices?|sauces?|cereals?|grains?)\b",
     gloss=r"\b(rice|flour|sugar|salt|oil|vinegar|soy sauce|fish sauce|paste|condiment|spice|seasoning|honey|syrup|jam)\b"),
 "Eggs & dairy": dict(cats=r"\b(dairy|eggs?)\b", gloss=r"\b(egg|milk|cheese|butter|cream|yogurt|dairy|curd)\b"),
 "Meat & poultry": dict(cats=r"\bmeats?\b", gloss=r"\b(meat|pork|beef|chicken|sausage|bacon|ham|poultry|duck|goat meat)\b"),
 "Pets": dict(cats=r"\b(dogs?|cats?|pets?)\b", gloss=r"\b(dog|cat|puppy|kitten|rabbit|hamster|parrot|pet|goldfish|turtle|tortoise)\b"),
 "Baked goods": dict(cats=r"\b(cakes?|breads?|pastr(y|ies)|rice cakes?)\b", gloss=r"\b(cake|bread|pastry|biscuit|cookie|rice cake|bun)\b"),
 "Beverages": dict(cats=r"\b(drinks?|beverages?)\b", gloss=r"\b(drink|beverage|juice|wine|beer|coffee|tea|liquor|cocoa)\b"),
 "Candy & sweets": dict(cats=r"\b(sweets?|candy|confection)\b", gloss=r"\b(candy|sweet|confection|brittle|caramel)\b"),
 "Snacks": dict(cats=r"\bsnacks?\b", gloss=r"\b(snack|fritter|chips?|crackers?)\b"),
 "Meals & fast food": dict(cats=r"\b(dishes|soups?|stews?)\b", gloss=r"\b(dish|soup|stew|noodles?|porridge|meal|viand)\b"),
 "Frozen treats": dict(cats=r"\b(ice cream|frozen)\b", gloss=r"\b(ice cream|popsicle|sherbet|shaved ice|frozen)\b"),
}
POOLS["Events & activities"] = POOLS.pop("Holiday/event")      # name used by the term bank
for _k, _v in DICT_RULES.items(): POOLS[_k]["dict"] = _v

# Hand-picked seeds (term=English meaning). These are common, everyday Filipino words/names/brands chosen to FIT
# a context, not to translate the English word. Tier = "Curated". Reviewer must still Approve each one.
# Person names: term=gender.  Currency amount pools are rule-based (see RULE_BASED).
CURATED = {
 "Places & buildings": ["tindahan=store", "sari-sari store=neighborhood variety store", "palengke=public market", "talipapa=small wet market",
   "karinderya=small eatery", "panaderya=bakery", "paaralan=school", "eskwelahan=school", "bahay=house", "bukid=farm", "hardin=garden",
   "bakuran=backyard", "parke=park", "plaza=town plaza", "simbahan=church", "kapilya=chapel", "ospital=hospital", "munisipyo=town hall",
   "barangay hall=barangay hall", "palaisdaan=fishpond", "taniman=planted field", "palayan=rice field", "dalampasigan=beach", "baryo=village",
   "kubo=nipa hut", "aklatan=library", "silid-aralan=classroom", "opisina=office", "gusali=building", "pamilihan=marketplace",
   "kainan=eatery", "bayan=town", "kapitbahayan=neighborhood"],
 "Containers & packaging": ["balde=pail", "timba=pail", "tabo=dipper", "bilao=round woven tray", "bayong=woven bag", "sako=sack", "lata=can",
   "garapon=jar", "bote=bottle", "kahon=box", "supot=bag", "banga=clay pot", "palanggana=basin", "kaldero=pot", "lalagyan=container",
   "tampipi=woven box", "buslo=basket", "bandehado=tray", "sachet=sachet", "tasa=cup", "baso=glass", "pitsel=pitcher", "dosena=dozen",
   "pares=pair", "pakete=pack"],
 "Household & furniture": ["lamesa=table", "silya=chair", "upuan=seat", "papag=bamboo bed", "banig=sleeping mat", "unan=pillow", "kumot=blanket",
   "kulambo=mosquito net", "aparador=cabinet", "duyan=hammock", "bangko=bench/stool", "tokador=dresser", "istante=shelf", "salamin=mirror",
   "walis=broom", "basahan=rag", "pandakot=dustpan", "bentilador=electric fan", "abaniko=hand fan", "plantsa=iron", "gripo=faucet",
   "kusina=kitchen", "kalan=stove", "sandok=ladle", "kawali=frying pan", "kutsara=spoon", "tinidor=fork", "pinggan=plate", "mangkok=bowl",
   "orasan=clock"],
 "Events & activities": ["pista=town fiesta", "Pasko=Christmas", "Bagong Taon=New Year", "Undas=All Saints' Day", "Noche Buena=Christmas Eve feast",
   "Simbang Gabi=dawn masses before Christmas", "Santacruzan=May procession", "Flores de Mayo=May flower festival", "binyag=baptism",
   "kasal=wedding", "kaarawan=birthday", "handaan=feast/party", "salu-salo=gathering", "pagtitipon=gathering", "paligsahan=contest",
   "konsiyerto=concert", "piknik=picnic", "bakasyon=vacation", "Araw ng Kalayaan=Independence Day", "Semana Santa=Holy Week",
   "Sinulog=Cebu festival", "Ati-Atihan=Aklan festival", "Pahiyas=Lucban festival", "Panagbenga=Baguio flower festival",
   "MassKara=Bacolod festival", "Kadayawan=Davao festival", "Dinagyang=Iloilo festival", "Buwan ng Wika=Filipino Language Month",
   "palaro=sports meet", "reunion=family reunion"],
 "Staples & ingredients": ["bigas=uncooked rice", "kanin=cooked rice", "asukal=sugar", "asin=salt", "harina=flour", "mantika=cooking oil",
   "suka=vinegar", "toyo=soy sauce", "patis=fish sauce", "bagoong=fermented shrimp/fish paste", "pulot=honey/syrup", "gata=coconut milk",
   "bawang=garlic", "sibuyas=onion", "luya=ginger", "paminta=pepper", "banana ketchup=banana ketchup", "sago=tapioca pearls",
   "asukal na pula=brown sugar", "mani=peanut"],
 "Eggs & dairy": ["itlog=egg", "itlog na maalat=salted egg", "balut=fertilized duck egg", "penoy=unfertilized duck egg", "itlog ng pugo=quail egg",
   "gatas=milk", "gatas ng kalabaw=carabao milk", "kesong puti=white carabao cheese", "keso=cheese", "mantikilya=butter", "krema=cream",
   "evaporada=evaporated milk", "kondensada=condensed milk"],
 "Meat & poultry": ["baboy=pork/pig", "manok=chicken", "baka=beef/cow", "kambing=goat", "itik=duck", "longganisa=sausage", "tocino=sweet cured pork",
   "tapa=cured beef", "liempo=pork belly", "lechon=roast pig", "hamon=ham", "karne=meat", "corned beef=corned beef", "chicharon=pork crackling",
   "pakpak ng manok=chicken wings", "atay=liver"],
 "Pets": ["aso=dog", "pusa=cat", "askal=mixed-breed dog", "puspin=mixed-breed cat", "kuneho=rabbit", "ibon=bird", "isda=fish", "pagong=turtle",
   "loro=parrot", "kalapati=pigeon", "tuta=puppy", "kuting=kitten", "alagang hayop=pet", "kanaryo=canary", "hamster=hamster"],
 "Aquatic animals": ["bangus=milkfish", "tilapia=tilapia", "galunggong=round scad", "tamban=sardine", "dilis=anchovy", "lapu-lapu=grouper", "hito=catfish",
   "dalag=mudfish", "alimasag=blue crab", "alimango=mud crab", "hipon=shrimp", "sugpo=tiger prawn", "pusit=squid", "pugita=octopus", "talaba=oyster",
   "tahong=mussel", "halaan=clam", "tulingan=bullet tuna", "pawikan=sea turtle", "butanding=whale shark", "pating=shark", "bisugo=threadfin bream",
   "maya-maya=snapper", "ayungin=silver perch"],
 "Vehicles & transport": ["jeepney=jeepney", "traysikel=tricycle", "kalesa=horse-drawn carriage", "bangka=outrigger boat", "habal-habal=motorcycle taxi",
   "padyak=pedicab", "kuliglig=motorized pedicab", "UV Express=van shuttle", "LRT=light rail", "MRT=metro rail", "bus=bus"],
 "Place/institution name": ["Rizal Elementary School=school", "Mabini High School=school", "Bonifacio Elementary School=school",
   "San Isidro Elementary School=school", "Sta. Cruz National High School=school", "Barangay San Roque Elementary School=school",
   "Tindahan ni Aling Nena=sari-sari store", "Karinderya ni Aling Rosa=eatery", "Panaderya ni Mang Berto=bakery",
   "Kainan ni Mang Tomas=eatery", "Aling Cora's Bakeshop=bakery", "Mang Pepe's Hardware=hardware store"],
 "Institution/government": ["pamahalaan=government", "barangay=smallest local government unit", "munisipyo=municipal hall", "DepEd=Department of Education",
   "PNP=national police", "LTO=Land Transportation Office", "BIR=Bureau of Internal Revenue", "SSS=Social Security System", "PhilHealth=health insurer",
   "Pag-IBIG=housing fund", "PBA=Philippine Basketball Association", "UAAP=university athletic association",
   "University of the Philippines=university", "Ateneo de Manila University=university", "De La Salle University=university"],
 "Street/road name": ["EDSA=Epifanio de los Santos Avenue", "Roxas Boulevard=Manila Bay road", "Rizal Avenue=Manila avenue", "Taft Avenue=Manila avenue",
   "Quezon Avenue=QC avenue", "Ayala Avenue=Makati avenue", "Commonwealth Avenue=QC avenue", "Katipunan Avenue=QC avenue",
   "Espana Boulevard=Manila boulevard", "Recto Avenue=Manila avenue", "Mabini Street=generic street", "Bonifacio Street=generic street",
   "Mabuhay Street=generic street"],
 "Place - country/region": ["Luzon=island group", "Visayas=island group", "Mindanao=island group", "Metro Manila=capital region", "Calabarzon=region",
   "Bicol Region=region", "Ilocos Region=region", "Cagayan Valley=region", "Central Luzon=region", "Western Visayas=region",
   "Central Visayas=region", "Davao Region=region", "Cordillera=region"],
 "Currency word": ["piso=peso", "pesos=pesos", "sentimo=centavo", "barya=loose change/coins", "singkwenta=fifty (centavos)"],
 "Brand/company": ["Jollibee=fast food", "Mang Inasal=grilled chicken chain", "Chowking=Chinese fast food", "Greenwich=pizza chain", "Max's=fried chicken restaurant",
   "Goldilocks=bakeshop", "Red Ribbon=bakeshop", "SM Supermarket=supermarket", "Puregold=supermarket", "Robinsons=mall/supermarket",
   "Mercury Drug=pharmacy", "Cebu Pacific=airline", "Globe=telecom", "Smart=telecom", "GCash=e-wallet", "Maya=e-wallet", "Shopee=online store",
   "Lazada=online store", "Skyflakes=crackers", "Piattos=chips", "Chippy=chips", "Nova=chips", "Boy Bawang=corn snack", "C2=iced tea",
   "Lucky Me=instant noodles", "Century Tuna=canned tuna", "Datu Puti=vinegar/soy sauce", "Choc Nut=chocolate candy", "Curly Tops=chocolate candy",
   "Royal=soda", "Sarsi=soda", "Magnolia=ice cream/dairy"],
 "Person name": ["Juan=male", "Jose=male", "Pedro=male", "Andres=male", "Emilio=male", "Ramon=male", "Antonio=male", "Manuel=male", "Carlos=male",
   "Miguel=male", "Rafael=male", "Gabriel=male", "Angelo=male", "Mark=male", "Paolo=male", "Jericho=male", "Renz=male", "Kenneth=male", "Christian=male",
   "Joshua=male", "Noel=male", "Rodel=male", "Rey=male", "Ronnie=male", "Arnel=male", "Dante=male", "Eduardo=male", "Fernando=male", "Francisco=male",
   "Ricardo=male", "Roberto=male", "Alfredo=male", "Benjamin=male", "Danilo=male", "Enrique=male", "Felipe=male", "Gerardo=male", "Isidro=male",
   "Jaime=male", "Mario=male", "Nestor=male", "Oscar=male", "Rolando=male", "Teodoro=male", "Vicente=male", "Wilfredo=male", "Jayson=male",
   "Carlo=male", "Luis=male", "Diego=male", "Marco=male", "Nico=male", "Rico=male", "Ryan=male", "Kevin=male", "Dodong=male",
   "Maria=female", "Ana=female", "Rosa=female", "Luz=female", "Carmen=female", "Josefa=female", "Teresa=female", "Elena=female", "Lourdes=female",
   "Cristina=female", "Angelica=female", "Jasmine=female", "Princess=female", "Kristine=female", "Kathleen=female", "Joanna=female", "Jessa=female",
   "Maricel=female", "Marites=female", "Liza=female", "Rowena=female", "Cherry=female", "Nena=female", "Remedios=female", "Perla=female",
   "Corazon=female", "Imelda=female", "Gloria=female", "Divina=female", "Estrella=female", "Glenda=female", "Hazel=female", "Irene=female",
   "Jocelyn=female", "Lorna=female", "Maribel=female", "Norma=female", "Pilar=female", "Rhea=female", "Sharon=female", "Tess=female", "Vilma=female",
   "Yolanda=female", "Zenaida=female", "Bea=female", "Sofia=female", "Isabella=female", "Gabriela=female", "Andrea=female", "Bianca=female",
   "Camille=female", "Erica=female", "Hannah=female", "Janine=female", "Kaye=female", "Mika=female", "Nicole=female", "Patricia=female", "Trisha=female"],
}

# Draft context attributes (keyword rules on description). A starting point for review, not ground truth.
ATTR_RULES = [
 ("countable pieces", r"\b(pieces?|rolls?|dumplings?|balls?|skewers?|buns?|cookies?|biscuits?|each|individual)\b"),
 ("sliceable / cut from a whole", r"\b(cake|pie|loaf|pudding|flan|slices?|sliced)\b"),
 ("liquid / drink", r"\b(drinks?|beverages?|juice|soup|broth|tea|coffee|milk|shake)\b"),
 ("sweet", r"\b(sweet|dessert|sugar|candy|syrup)\b"),
 ("fried", r"\b(fried|deep-fried|crispy)\b"),
 ("street food", r"\b(street food|street vendor|skewer)\b"),
 ("rice-based", r"\b(rice|glutinous|kakanin)\b"),
 ("rideable / vehicle", r"\b(vehicle|jeepney|tricycle|bus|boat|ferry|taxi|bicycle|train)\b"),
 ("wearable", r"\b(clothing|garment|dress|shirt|hat|sandals?|footwear|attire)\b"),
 ("game / toy", r"\b(game|toy|played|players?)\b"),
 ("animal", r"\b(species|mammal|bird|reptile|fish|insect|endemic)\b"),
 ("plant", r"\b(tree|plant|flower|shrub|fruit)\b"),
]

LICENSE = "Wikipedia/Wiktionary text: CC BY-SA 4.0; Wikidata: CC0"

# --------------------------------------------------------------------------------------
# HTTP with on-disk cache, retries and polite pacing
# --------------------------------------------------------------------------------------
class Client:
    def __init__(self, cache_dir, contact, delay=0.15):
        self.cache = Path(cache_dir); self.cache.mkdir(parents=True, exist_ok=True)
        self.ua = f"term-pool-extractor/1.0 ({contact})"
        self.delay = delay
        self.s = requests.Session(); self.s.headers["User-Agent"] = self.ua

    def _request(self, url, params):
        """One network call. Returns (status_code, json_or_None). Overridden in tests."""
        for attempt in range(5):
            r = self.s.get(url, params=params, timeout=30)
            if r.status_code in (429, 500, 502, 503, 504):
                time.sleep(float(r.headers.get("Retry-After", 2 ** attempt))); continue
            if r.status_code == 404: return 404, None
            r.raise_for_status()
            return r.status_code, r.json()
        raise RuntimeError(f"Giving up on {url}")

    def get(self, url, params=None):
        params = params or {}
        key = hashlib.sha1((url + json.dumps(params, sort_keys=True)).encode()).hexdigest()
        f = self.cache / f"{key}.json"
        if f.exists():
            return json.loads(f.read_text(encoding="utf-8"))
        time.sleep(self.delay)
        status, data = self._request(url, params)
        if status == 404: data = None
        f.write_text(json.dumps(data), encoding="utf-8")
        return data

def api(host): return f"https://{host}/w/api.php"
def project(host): return host.replace(".org", "")

# --------------------------------------------------------------------------------------
# Source readers
# --------------------------------------------------------------------------------------
def category_titles(cl, host, category, depth, cap):
    """Main-namespace page titles in a category tree (BFS). Returns (titles, found_any)."""
    titles, seen, queue = [], set(), [(category, 0)]
    while queue and len(titles) < cap:
        cat, d = queue.pop(0)
        if cat in seen: continue
        seen.add(cat)
        cont = {}
        while True:
            p = dict(action="query", list="categorymembers", cmtitle="Category:" + cat, cmlimit=500,
                     cmtype="page|subcat", format="json", formatversion=2, **cont)
            data = cl.get(api(host), p) or {}
            for m in data.get("query", {}).get("categorymembers", []):
                if m.get("ns") == 0: titles.append(m["title"])
                elif m.get("ns") == 14 and d < depth: queue.append((m["title"].split(":", 1)[1], d + 1))
            if "continue" in data: cont = data["continue"]
            else: break
    return list(dict.fromkeys(titles))[:cap]

def batched(seq, n):
    seq = list(seq)
    for i in range(0, len(seq), n): yield seq[i:i + n]

def page_info(cl, host, titles):
    """title -> dict(title, extract, qid) with redirects resolved."""
    out = {}
    for batch in batched(titles, 20):
        cont = {}
        while True:
            p = dict(action="query", prop="extracts|pageprops", exintro=1, explaintext=1, exsentences=2,
                     exlimit=20, ppprop="wikibase_item", redirects=1, titles="|".join(batch),
                     format="json", formatversion=2, **cont)
            data = cl.get(api(host), p) or {}
            for pg in data.get("query", {}).get("pages", []):
                if pg.get("missing"): continue
                rec = out.setdefault(pg["title"], dict(title=pg["title"], extract="", qid=None))
                if pg.get("extract"): rec["extract"] = pg["extract"].strip()
                q = (pg.get("pageprops") or {}).get("wikibase_item")
                if q: rec["qid"] = q
            if "continue" in data: cont = data["continue"]
            else: break
    return out

def wikidata(cl, qids):
    """qid -> dict(en, tl, desc, aliases_tl, tl_title, instance_of[], origin_ph)."""
    ents = {}
    for batch in batched(sorted(set(qids)), 50):
        p = dict(action="wbgetentities", ids="|".join(batch), props="labels|descriptions|aliases|sitelinks|claims",
                 languages="en|tl", sitefilter="tlwiki", format="json")
        data = cl.get("https://www.wikidata.org/w/api.php", p) or {}
        for qid, e in (data.get("entities") or {}).items():
            if "missing" in e: continue
            claims = e.get("claims", {})
            def ids(prop):
                out = []
                for c in claims.get(prop, []):
                    v = (c.get("mainsnak", {}).get("datavalue") or {}).get("value")
                    if isinstance(v, dict) and v.get("id"): out.append(v["id"])
                return out
            ents[qid] = dict(
                en=(e.get("labels", {}).get("en") or {}).get("value"),
                tl=(e.get("labels", {}).get("tl") or {}).get("value"),
                desc=(e.get("descriptions", {}).get("en") or {}).get("value", ""),
                aliases_tl=[a["value"] for a in e.get("aliases", {}).get("tl", [])],
                tl_title=(e.get("sitelinks", {}).get("tlwiki") or {}).get("title"),
                p31=ids("P31"), origin_ph="Q928" in ids("P495"))  # Q928 = Philippines
    # resolve instance-of labels
    classes = sorted({c for e in ents.values() for c in e["p31"]})
    labels = {}
    for batch in batched(classes, 50):
        p = dict(action="wbgetentities", ids="|".join(batch), props="labels", languages="en", format="json")
        data = cl.get("https://www.wikidata.org/w/api.php", p) or {}
        for qid, e in (data.get("entities") or {}).items():
            labels[qid] = (e.get("labels", {}).get("en") or {}).get("value", qid)
    for e in ents.values(): e["instance_of"] = [labels.get(c, c) for c in e["p31"]]
    return ents

def month_window():
    t = dt.date.today().replace(day=1)
    end = (t - dt.timedelta(days=1)).replace(day=1)
    y, m = end.year, end.month - 11
    while m <= 0: m += 12; y -= 1
    return dt.date(y, m, 1).strftime("%Y%m%d"), end.strftime("%Y%m%d")

def pageviews(cl, proj, title):
    start, end = month_window()
    url = (f"https://wikimedia.org/api/rest_v1/metrics/pageviews/per-article/{proj}/all-access/user/"
           f"{quote(title.replace(' ', '_'), safe='')}/monthly/{start}/{end}")
    data = cl.get(url)
    return sum(i.get("views", 0) for i in (data or {}).get("items", []))

# --------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------
def clean_title(t): return re.sub(r"\s*\([^)]*\)\s*$", "", t).strip()
def norm(t):
    t = unicodedata.normalize("NFKD", t.lower())
    return re.sub(r"\s+", " ", "".join(c for c in t if not unicodedata.combining(c))).strip()
def skip_title(t):
    return (t.startswith(("List of", "Lists of", "Index of", "Outline of")) or ":" in t
            or "disambiguation" in t.lower())
def attrs_for(text):
    return ", ".join(name for name, rx in ATTR_RULES if re.search(rx, text, re.I))

# --------------------------------------------------------------------------------------
# Pipeline
# --------------------------------------------------------------------------------------
def build_pool(cl, pool, cfg, cap, with_views, log, src_rows, workers):
    must = re.compile(r"\b(?:" + cfg["must"] + r")", re.I) if cfg.get("must") else None
    ex_pat = cfg.get("exclude", DEFAULT_EXCLUDE) if "exclude" in cfg else DEFAULT_EXCLUDE
    exclude = re.compile(ex_pat, re.I) if ex_pat else None
    cands = {}  # norm term -> row
    for spec in cfg["sources"]:
        host, cat = spec["host"], spec["category"]
        titles = category_titles(cl, host, cat, spec["depth"], cap)
        found = len(titles)
        if found == 0:
            log.append((pool, f"{host} Category:{cat}", "0 members (category missing or empty): run `discover` and fix POOLS"))
        titles = [t for t in titles if not skip_title(t)]
        info = page_info(cl, host, titles)
        qids = [i["qid"] for i in info.values() if i["qid"]]
        wd = wikidata(cl, qids) if qids else {}
        kept = 0
        for t, i in info.items():
            w = wd.get(i["qid"], {}) if i["qid"] else {}
            text = " ".join([t, w.get("desc", ""), i["extract"]])
            first = i["extract"].split(".")[0] if i["extract"] else ""
            if must and not must.search(text): continue
            if exclude and exclude.search(first + " " + w.get("desc", "")): continue
            en = clean_title(w.get("en") or t)
            tl = clean_title(w.get("tl") or w.get("tl_title") or "") or None
            use = tl or en
            key = norm(use)
            url = f"https://{host}/wiki/{quote(t.replace(' ', '_'))}"
            row = dict(pool=pool, term=use, english=en, tagalog=tl or "", aliases=", ".join(w.get("aliases_tl", [])),
                       desc=(w.get("desc") or first)[:300], attrs=attrs_for(text),
                       instance_of=", ".join(w.get("instance_of", [])[:4]), origin_ph="Yes" if w.get("origin_ph") else "",
                       qid=i["qid"] or "", url=url, title=t, host=host, tl_title=w.get("tl_title"),
                       source=f"{host} > Category:{cat}", views_en=0, views_tl=0)
            if key in cands:
                cands[key]["source"] += f"; {row['source']}"
            else:
                cands[key] = row
            kept += 1
        src_rows.append((pool, host, cat, spec["depth"], found, kept))
    rows = list(cands.values())
    if with_views and rows:
        def job(r):
            r["views_en"] = pageviews(cl, project(r["host"]), r["title"])
            if r["tl_title"]: r["views_tl"] = pageviews(cl, "tl.wikipedia", r["tl_title"])
        with ThreadPoolExecutor(max_workers=workers) as ex: list(ex.map(job, rows))
    for r in rows:
        r["score"] = round(math.log10(1 + r["views_en"]) + 1.5 * math.log10(1 + r["views_tl"])
                           + (0.5 if r["origin_ph"] else 0), 3)
    ranked = sorted(rows, key=lambda r: -r["score"])
    n = len(ranked)
    for idx, r in enumerate(ranked):
        r["tier"] = ("n/a (views skipped)" if not with_views else
                     "High" if idx < n / 3 else "Medium" if idx < 2 * n / 3 else "Low")
    return ranked

# --------------------------------------------------------------------------------------
# NEW: generic-term dictionary source, curated seeds, names file, term-bank targets
# --------------------------------------------------------------------------------------
KAIKKI_URL = "https://kaikki.org/dictionary/Tagalog/kaikki.org-dictionary-Tagalog.jsonl"   # verify on kaikki.org/dictionary/Tagalog/
BAD_TAGS = {"archaic", "obsolete", "dialectal", "rare", "dated", "historical", "vulgar", "offensive", "derogatory",
            "slang", "poetic", "nonstandard", "misspelling", "abbreviation", "initialism", "uncommon", "colloquial"}
LOAN_RX = re.compile(r"(borrowed|borrowings|derived) from English", re.I)
ACTIONS = ("Localize", "Localize (check context)")
TERMBANK_SHEET = "Lookup Table (Cleaned)"

def _open_text(path):
    p = str(path)
    return gzip.open(p, "rt", encoding="utf-8") if p.endswith(".gz") else open(p, encoding="utf-8")

def _cat_names(cats):
    return [c.get("name", "") if isinstance(c, dict) else str(c) for c in cats or []]

def load_dictionary(path, allow_english_loans=False):
    """Usable Tagalog noun entries from a kaikki.org (wiktextract) JSONL file.
    Drops: non-nouns, phrases of 3+ words, affixes, archaic/dialectal/rare/slang senses, alt-spellings, and (by default)
    words Wiktionary marks as borrowed from English (key, kompyuter...): those are not 'Filipino substitutes'.
    Spanish loans (mesa, silya, kutsara) are kept: they are ordinary Filipino."""
    out = []
    with _open_text(path) as f:
        for line in f:
            try: e = json.loads(line)
            except ValueError: continue
            if e.get("lang_code") != "tl" or e.get("pos") != "noun": continue
            w = (e.get("word") or "").strip()
            if not w or len(w.split()) > 2 or not re.fullmatch(r"[\w' -]+", w) or w[0] in "-'": continue
            cats = _cat_names(e.get("categories"))
            loan = any(LOAN_RX.search(c) for c in cats)
            if loan and not allow_english_loans: continue
            glosses, topics = [], []
            for s in e.get("senses", []):
                if set(s.get("tags", [])) & BAD_TAGS or s.get("form_of") or s.get("alt_of"): continue
                glosses += [g for g in s.get("glosses", []) if g]
                topics += s.get("topics", [])
                cats += _cat_names(s.get("categories"))
            if glosses:
                out.append(dict(word=w, glosses=glosses, loan=loan, text=" | ".join(dict.fromkeys(cats + topics)).lower()))
    return out

ILLEGAL_RX = re.compile(r"[\000-\010\013\014\016-\037]")      # characters openpyxl refuses in cells

def _blank_row(pool, term, english, desc, source, **extra):
    row = dict(pool=pool, term=term, english=english, tagalog=term, aliases="", desc=desc[:300], attrs=attrs_for(desc),
               instance_of="", origin_ph="", qid="", url="", title=term, host="", tl_title=None, source=source,
               views_en=0, views_tl=0, score=0, tier="")
    row.update(extra)
    for k in ("term", "english", "tagalog", "desc"): row[k] = ILLEGAL_RX.sub("", row[k]) if isinstance(row[k], str) else row[k]
    return row

def dict_rows(pool, cfg, entries, cap):
    dc = cfg.get("dict")
    if not dc: return []
    cat_rx = re.compile(dc["cats"], re.I) if dc.get("cats") else None
    gl_rx = re.compile(dc["gloss"], re.I) if dc.get("gloss") else None
    hits = []
    for e in entries:
        by_cat = bool(cat_rx and cat_rx.search(e["text"]))
        by_gl = bool(gl_rx and gl_rx.search(" ; ".join(e["glosses"][:2])))
        if by_cat or by_gl: hits.append((-(by_cat + by_gl), len(e["word"].split()), len(e["word"]), e, by_cat))
    hits.sort(key=lambda h: h[:3])          # matched by both category and gloss first, then single words, then short words
    rows = []
    for _, _, _, e, by_cat in hits[:cap]:
        w = e["word"]
        r = _blank_row(pool, w, e["glosses"][0][:120], "; ".join(e["glosses"][:3]),
                       "kaikki.org Tagalog Wiktionary (" + ("category" if by_cat else "gloss") + " match)",
                       url=f"https://en.wiktionary.org/wiki/{quote(w.replace(' ', '_'))}#Tagalog", host="en.wiktionary.org",
                       tier="Dictionary (unrated)", dictionary=True)
        rows.append(r)
    return rows

def curated_rows(pool, items):
    rows = []
    for it in items:
        term, _, eng = it.partition("=")
        rows.append(_blank_row(pool, term.strip(), eng.strip(), eng.strip() or term.strip(),
                               "Curated seed (CURATED in extract_pool.py) - review", tier="Curated", curated=True))
    return rows

def names_rows(path):
    """CSV/TXT with  name[,count[,gender]]  -> rows. A header row (name/first name/given name) is skipped."""
    raw = []
    data = Path(path).read_bytes()          # encoding varies by how the file was saved (Excel = cp1252, PowerShell > = UTF-16)
    if data[:4] == b"PK\x03\x04":            # an Excel workbook saved/renamed with a .csv extension
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        recs = [["" if c is None else str(c) for c in row] for row in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close(); print(f"NOTE: {path} is an Excel workbook, reading its first sheet (columns: name, count, gender)", flush=True)
    else:
        if data[:2] in (b"\xff\xfe", b"\xfe\xff"): text = data.decode("utf-16")
        else:
            try: text = data.decode("utf-8-sig")
            except UnicodeDecodeError: text = data.decode("cp1252", errors="replace")
        recs = list(csv.reader(text.splitlines()))
    skipped = 0
    for rec in recs:
        if not rec or not rec[0].strip() or rec[0].strip().lower() in ("name", "first name", "given name", "firstname"): continue
        name = rec[0].strip()
        if len(name) > 60 or not name.isprintable(): skipped += 1; continue      # binary/garbage row
        try: cnt = float(rec[1].replace(",", "")) if len(rec) > 1 and rec[1].strip() else 0
        except ValueError: cnt = 0
        raw.append((name, cnt, rec[2].strip() if len(rec) > 2 else ""))
    if skipped: print(f"WARNING: skipped {skipped} unreadable rows in {path}", flush=True)
    raw.sort(key=lambda x: -x[1]); n = len(raw); rows = []
    for i, (name, cnt, gender) in enumerate(raw):
        rows.append(_blank_row("Person name", name, gender, f"{gender} given name".strip(), f"names file: {Path(path).name}",
                               views_en=int(cnt), score=round(math.log10(1 + cnt), 3),
                               tier=("High" if i < n / 3 else "Medium" if i < 2 * n / 3 else "Low") if cnt else "n/a (no counts)"))
    return rows

def merge_extra(pool, cfg, rows, entries, args, log, src_rows):
    """Add curated seeds, dictionary matches and the names file to the Wikipedia rows (dedupe by normalised term)."""
    seen = {norm(r["term"]): r for r in rows}
    def add(new, label, host, where):
        added = 0
        for r in new:
            k = norm(r["term"])
            if k in seen:
                if label not in seen[k]["source"]: seen[k]["source"] += f"; {label}"
                if r.get("curated"): seen[k]["curated"] = True; seen[k]["tier"] = "Curated"
                if r.get("views_en") and not seen[k].get("views_en"): seen[k]["views_en"] = r["views_en"]; seen[k]["score"] = r["score"]   # keep names-file counts
                continue
            seen[k] = r; rows.append(r); added += 1
        src_rows.append((pool, host, where, "-", len(new), added))
    if CURATED.get(pool): add(curated_rows(pool, CURATED[pool]), "curated seed", "curated", "CURATED list")
    if cfg.get("dict"):
        if entries is None: log.append((pool, "dictionary", "No --dict file given: Tagalog dictionary source skipped for this pool"))
        else: add(dict_rows(pool, cfg, entries, args.dict_cap), "kaikki Tagalog Wiktionary", "kaikki.org", "dictionary match (cap %d)" % args.dict_cap)
    if pool == "Person name" and args.names_file and Path(args.names_file).exists():
        add(names_rows(args.names_file), "names file", "names file", Path(args.names_file).name)
    elif pool == "Person name" and args.names_file:
        log.append((pool, "names", f"--names-file {args.names_file} not found: skipped (only Wikipedia + CURATED names used)"))
        print(f"WARNING: names file {args.names_file} not found: skipped", flush=True)
    elif pool == "Person name": log.append((pool, "names", "No --names-file: only Wikipedia + CURATED names. The term bank needs ~2,400 names, so expect a large shortfall"))
    return rows

def split_recognized(rows, cfg, args, with_views):
    """Wikipedia rows below --min-score are moved to the 'Rejected (auto)' sheet (kept for audit, not silently lost)."""
    if not with_views or args.min_score <= 0 or cfg.get("rec") is False: return rows, []
    keep, rej = [], []
    for r in rows:
        if r.get("curated") or r.get("dictionary") or r["score"] >= args.min_score: keep.append(r)
        else: r["reject_reason"] = f"recognition score {r['score']} < {args.min_score}"; rej.append(r)
    return keep, rej

def load_targets(path):
    """Term bank -> ({pool: {firm, ctx, freq}}, [(term, freq, pool, action)]) for Localize / Localize (check context) rows."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    it = wb[TERMBANK_SHEET].iter_rows(values_only=True)
    ix = {h: i for i, h in enumerate(next(it))}
    pools, terms = {}, []
    for row in it:
        act, pool = row[ix["Suggested Action"]], row[ix["Replacement Pool"]]
        if act not in ACTIONS or not pool: continue
        freq = int(row[ix["Frequency"]] or 0)
        p = pools.setdefault(pool, dict(firm=0, ctx=0, freq=0))
        p["firm" if act == ACTIONS[0] else "ctx"] += 1; p["freq"] += freq
        terms.append((row[ix["Term"]], freq, pool, act))
    wb.close(); return pools, terms

def literal_index(entries):
    idx = {}
    for e in entries:
        for g in e["glosses"][:3]:
            for part in re.split(r"[;,]", g):
                k = re.sub(r"^(a|an|the|to) ", "", norm(re.sub(r"\([^)]*\)", "", part)))
                if k: idx.setdefault(k, []).append(e["word"])
    return idx

def build_generic(terms, entries, chosen):
    """Tagged generic terms with their LITERAL dictionary translation, so reviewers can see what NOT to use blindly."""
    idx = literal_index(entries); rows = []
    for term, freq, pool, act in sorted(terms, key=lambda t: -t[1]):
        if pool not in chosen or not POOLS.get(pool, {}).get("dict"): continue
        k = norm(term); hits = idx.get(k) or idx.get(k.rstrip("s")) or []
        rows.append((term, freq, pool, act, ", ".join(dict.fromkeys(hits))[:120]))
    return rows

# --------------------------------------------------------------------------------------
# Output
# --------------------------------------------------------------------------------------
COLS = [("Pool", "pool", 26), ("Term (use)", "term", 26), ("English", "english", 26), ("Tagalog", "tagalog", 22),
        ("Tagalog aliases", "aliases", 24), ("Description", "desc", 60), ("Context attributes (draft)", "attrs", 30),
        ("Recognition tier", "tier", 16), ("Score", "score", 8), ("Views EN (12 mo)", "views_en", 14),
        ("Views TL (12 mo)", "views_tl", 14), ("PH origin (Wikidata)", "origin_ph", 12), ("Instance of", "instance_of", 28),
        ("Wikidata ID", "qid", 12), ("Page URL", "url", 44), ("Source", "source", 44)]

def write_xlsx(path, rows, src_rows, log, args, targets=None, rejected=(), generic=()):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    F = "Arial"; hf = Font(name=F, bold=True, color="FFFFFF", size=10); fill = PatternFill("solid", fgColor="1F3864")
    bf = Font(name=F, size=10)
    wb = Workbook(); ws = wb.active; ws.title = "README"
    text = [
        ("Replacement pool - extracted terms", True), ("", False),
        (f"Generated {dt.datetime.now():%Y-%m-%d %H:%M}. Pools: {', '.join(sorted({r['pool'] for r in rows}))}.", False), ("", False),
        ("Sources", True),
        ("Wikipedia/Wiktionary category trees (MediaWiki API) for candidate terms and short descriptions.", False),
        ("Wikidata for Tagalog labels, aliases, instance-of and country of origin (Philippines).", False),
        ("Wikimedia Pageviews API for 12-month views, used as a recognizability score.", False),
        (f"Licensing: {LICENSE}. Keep the Source and Page URL columns when redistributing.", False), ("", False),
        ("How to read it", True),
        ("Recognition tier = rank of Score within the pool (High = top third). Score = log10(EN views) + 1.5 x log10(TL views) + 0.5 if Wikidata says PH origin.", False),
        ("Context attributes are keyword guesses from the description (countable, sliceable, liquid, ...). Review them before use.", False),
        ("Reviewer status / notes columns are empty on purpose: set Approve / Reject / Maybe per term.", False), ("", False),
        ("Known limits", True),
        ("Wikipedia coverage favors Manila-centric and well-documented items; regional terms may be missing.", False),
        ("Category seeds are guesses until verified: see the Run Log sheet for categories that returned nothing.", False), ("", False),
        ("Generic nouns, curated seeds, coverage", True),
        ("Generic nouns (places, containers, household, pets, ...) come from the kaikki.org Tagalog Wiktionary extract (--dict) plus the CURATED seed lists. They are context-fitting candidates, not equivalents: Approve what fits.", False),
        ("Native-word filter: entries Wiktionary marks as borrowed from English, and archaic/dialectal/rare/slang senses, are dropped from the dictionary source. Spanish loans (silya, kutsara) are kept.", False),
        ("Coverage compares each pool with the term bank (Localize + check-context unique terms, plus the buffer). Currency amount pools are rule-based, not term pools.", False),
        ("Rejected (auto) lists Wikipedia candidates below --min-score recognition, so nothing is silently lost. Generic Terms shows the literal dictionary translation of each tagged term (what NOT to use blindly).", False),
        ("Wiktionary data via kaikki.org: cite Ylonen, Wiktextract, LREC 2022. CC BY-SA 4.0.", False),
    ]
    for i, (t, b) in enumerate(text, 1):
        c = ws.cell(row=i, column=1, value=t); c.font = Font(name=F, size=12 if i == 1 else 10, bold=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130

    pw = wb.create_sheet("Pool")
    heads = [h for h, _, _ in COLS] + ["Reviewer status", "Reviewer notes"]
    for j, h in enumerate(heads, 1):
        c = pw.cell(row=1, column=j, value=h); c.font = hf; c.fill = fill
    for i, r in enumerate(rows, 2):
        for j, (_, k, _) in enumerate(COLS, 1):
            pw.cell(row=i, column=j, value=r.get(k)).font = bf
    for j, (_, _, w) in enumerate(COLS, 1): pw.column_dimensions[get_column_letter(j)].width = w
    for j in (len(COLS) + 1, len(COLS) + 2): pw.column_dimensions[get_column_letter(j)].width = 18
    pw.freeze_panes = "C2"; pw.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{max(len(rows)+1, 2)}"
    dv = DataValidation(type="list", formula1='"Approve,Reject,Maybe"', allow_blank=True)
    pw.add_data_validation(dv); col = get_column_letter(len(COLS) + 1); dv.add(f"{col}2:{col}{max(len(rows)+1, 2)}")

    sw = wb.create_sheet("Sources")
    for j, h in enumerate(["Pool", "Host", "Category", "Depth", "Titles found", "Rows kept after filters"], 1):
        c = sw.cell(row=1, column=j, value=h); c.font = hf; c.fill = fill
    for i, row in enumerate(src_rows, 2):
        for j, v in enumerate(row, 1): sw.cell(row=i, column=j, value=v).font = bf
    for j, w in enumerate([28, 22, 44, 8, 14, 22], 1): sw.column_dimensions[get_column_letter(j)].width = w

    lw = wb.create_sheet("Run Log")
    for j, h in enumerate(["Pool", "Where", "Message"], 1):
        c = lw.cell(row=1, column=j, value=h); c.font = hf; c.fill = fill
    for i, row in enumerate(log or [("-", "-", "No warnings")], 2):
        for j, v in enumerate(row, 1): lw.cell(row=i, column=j, value=v).font = bf
    for j, w in enumerate([28, 52, 90], 1): lw.column_dimensions[get_column_letter(j)].width = w
    if targets:
        cw = wb.create_sheet("Coverage", 1)
        cw["A1"] = "Excess buffer over tagged unique terms"; cw["A1"].font = Font(name=F, bold=True, size=10)
        cw["B1"] = args.buffer; cw["B1"].number_format = "0%"; cw["B1"].font = Font(name=F, size=10, color="0000FF")
        cw["C1"] = "Input (blue): change it and every Target recalculates."; cw["C1"].font = Font(name=F, size=9, italic=True)
        heads2 = ["Pool", "Tagged unique: Localize", "Tagged unique: check context", "Tagged unique: total", "Target (with buffer)",
                  "Tagged occurrences", "Pool candidates", "Candidates not rejected", "Shortfall vs target", "Status"]
        for j, h in enumerate(heads2, 1):
            c = cw.cell(row=3, column=j, value=h); c.font = hf; c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical="center")
        stat = get_column_letter(len(COLS) + 1)       # Reviewer status column on the Pool sheet
        order = sorted(targets, key=lambda p: -(targets[p]["firm"] + targets[p]["ctx"]))
        for i, p in enumerate(order, 4):
            t = targets[p]
            cw.cell(row=i, column=1, value=p); cw.cell(row=i, column=2, value=t["firm"]); cw.cell(row=i, column=3, value=t["ctx"])
            cw.cell(row=i, column=4, value=f"=B{i}+C{i}")
            cw.cell(row=i, column=5, value=f"=ROUNDUP(D{i}*(1+$B$1),0)"); cw.cell(row=i, column=6, value=t["freq"])
            if p in RULE_BASED:
                for col in (7, 8, 9): cw.cell(row=i, column=col, value="n/a")
                cw.cell(row=i, column=10, value="Rule-based (convert amounts, not a term pool)")
            else:
                cw.cell(row=i, column=7, value=f"=COUNTIF(Pool!$A:$A,A{i})")
                cw.cell(row=i, column=8, value=f'=COUNTIFS(Pool!$A:$A,A{i},Pool!${stat}:${stat},"<>Reject")')
                cw.cell(row=i, column=9, value=f"=MAX(0,E{i}-H{i})")
                cw.cell(row=i, column=10, value=f'=IF(I{i}=0,"OK","SHORT")')
        last = 3 + len(order); tot = last + 1
        cw.cell(row=tot, column=1, value="Total (term pools only)")
        for col in (2, 3, 4, 5, 6):
            L = get_column_letter(col); cw.cell(row=tot, column=col, value=f"=SUM({L}4:{L}{last})")
        for col in (7, 8, 9):
            L = get_column_letter(col); cw.cell(row=tot, column=col, value=f"=SUM({L}4:{L}{last})")
        for r_ in cw.iter_rows(min_row=4, max_row=tot):
            for c in r_: c.font = Font(name=F, size=10, bold=(c.row == tot))
        cw.cell(row=tot + 2, column=1, value="Target = ROUNDUP((Localize + check-context unique terms) x (1 + buffer)). Many 'check context' terms (box, party, milk, dog) may end up Keep, so the real need can be lower: shortfall on firm Localize terms is the hard floor.").font = Font(name=F, size=9, italic=True)
        cw.cell(row=tot + 3, column=1, value="Candidates not rejected counts every Pool row whose Reviewer status is not Reject (blank = still pending). Only pools run in this workbook have rows.").font = Font(name=F, size=9, italic=True)
        for j, w in enumerate([36, 14, 14, 14, 14, 14, 14, 14, 14, 40], 1): cw.column_dimensions[get_column_letter(j)].width = w
        cw.row_dimensions[3].height = 32; cw.freeze_panes = "B4"
    if generic:
        gw = wb.create_sheet("Generic Terms")
        for j, h in enumerate(["Tagged term", "Occurrences", "Pool", "Suggested action", "Literal Tagalog (dictionary)", "Note"], 1):
            c = gw.cell(row=1, column=j, value=h); c.font = hf; c.fill = fill
        for i, (term, freq, pool, act, lit) in enumerate(generic, 2):
            for j, v in enumerate([term, freq, pool, act, lit, "literal match: review before use; prefer a context-fitting pool term" if lit else "no literal match"], 1):
                gw.cell(row=i, column=j, value=v).font = bf
        for j, w in enumerate([26, 12, 28, 24, 36, 60], 1): gw.column_dimensions[get_column_letter(j)].width = w
        gw.freeze_panes = "A2"; gw.auto_filter.ref = f"A1:F{len(generic)+1}"
    if rejected:
        rw = wb.create_sheet("Rejected (auto)")
        for j, h in enumerate(["Pool", "Term", "English", "Score", "Views EN (12 mo)", "Views TL (12 mo)", "Reason", "Page URL"], 1):
            c = rw.cell(row=1, column=j, value=h); c.font = hf; c.fill = fill
        for i, r in enumerate(rejected, 2):
            for j, v in enumerate([r["pool"], r["term"], r["english"], r["score"], r["views_en"], r["views_tl"], r.get("reject_reason", ""), r["url"]], 1):
                rw.cell(row=i, column=j, value=v).font = bf
        for j, w in enumerate([26, 28, 28, 8, 14, 14, 36, 50], 1): rw.column_dimensions[get_column_letter(j)].width = w
        rw.freeze_panes = "A2"; rw.auto_filter.ref = f"A1:H{len(rejected)+1}"
    wb.save(path)

def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow([h for h, _, _ in COLS])
        for r in rows: w.writerow([r.get(k, "") for _, k, _ in COLS])

# --------------------------------------------------------------------------------------
# Commands
# --------------------------------------------------------------------------------------
def cmd_run(args):
    cl = Client(args.cache, args.contact, args.delay)
    chosen = [p.strip() for p in args.pools.split(",")] if args.pools else list(POOLS)
    unknown = [p for p in chosen if p not in POOLS]
    if unknown: sys.exit(f"Unknown pools: {unknown}. Use list-pools.")
    targets = terms = None
    if args.targets and Path(args.targets).exists():
        targets, terms = load_targets(args.targets)
        print(f"[term bank] {sum(t['firm'] + t['ctx'] for t in targets.values())} tagged unique terms in {len(targets)} pools")
    entries = None
    if args.dict and not Path(args.dict).exists():
        sys.exit(f"Dictionary file {args.dict} not found. Run fetch-dict first, or drop --dict.")
    if args.dict:
        entries = load_dictionary(args.dict, args.allow_english_loans)
        print(f"[dictionary] {len(entries)} usable Tagalog noun entries")
    all_rows, src_rows, log, rejected = [], [], [], []
    for pool in chosen:
        cfg = POOLS[pool]
        print(f"[{pool}] ...", flush=True)
        rows = build_pool(cl, pool, cfg, args.max_per_source, not args.no_views, log, src_rows, args.workers)
        rows = merge_extra(pool, cfg, rows, entries, args, log, src_rows)
        rows, rej = split_recognized(rows, cfg, args, not args.no_views)
        rejected += rej
        print(f"[{pool}] {len(rows)} terms, {len(rej)} auto-rejected (low recognition)", flush=True)
        all_rows += rows
    if targets:
        for p in sorted(targets):
            if p not in POOLS and p not in RULE_BASED:
                log.append((p, "coverage", "Term-bank pool has no source in POOLS/CURATED: add one"))
    generic = build_generic(terms, entries, set(chosen)) if terms and entries else []
    write_xlsx(args.out, all_rows, src_rows, log, args, targets, rejected, generic)
    if args.csv: write_csv(args.csv, all_rows)
    print(f"Wrote {args.out} ({len(all_rows)} rows, {len(rejected)} auto-rejected); {len(log)} warnings (see Run Log sheet).")

def cmd_targets(args):
    pools, _ = load_targets(args.targets)
    rows = []
    for p, t in sorted(pools.items(), key=lambda kv: -(kv[1]["firm"] + kv[1]["ctx"])):
        tot = t["firm"] + t["ctx"]
        rows.append((p, t["firm"], t["ctx"], tot, math.ceil(tot * (1 + args.buffer)), math.ceil(t["firm"] * (1 + args.buffer)), t["freq"],
                     "rule-based" if p in RULE_BASED else ("dict+curated" if POOLS.get(p, {}).get("dict") else "wikipedia/curated" if p in POOLS else "NO SOURCE")))
    print(f"{'pool':42s} {'firm':>6s} {'ctx':>6s} {'total':>6s} {'target':>7s} {'t(firm)':>7s} {'occur':>7s}  source")
    for r in rows: print(f"{r[0]:42s} {r[1]:6d} {r[2]:6d} {r[3]:6d} {r[4]:7d} {r[5]:7d} {r[6]:7d}  {r[7]}")
    if args.csv:
        with open(args.csv, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f); w.writerow(["pool", "localize", "check_context", "total", "target", "target_localize_only", "occurrences", "source"]); w.writerows(rows)

def cmd_fetch_dict(args):
    r = requests.get(args.url, stream=True, timeout=60, headers={"User-Agent": f"term-pool-extractor/1.0 ({args.contact})"})
    r.raise_for_status(); n = 0
    with open(args.out, "wb") as f:
        for chunk in r.iter_content(1 << 20): f.write(chunk); n += len(chunk)
    print(f"Wrote {args.out} ({n / 1e6:.1f} MB) from {args.url}")

def cmd_discover_dict(args):
    from collections import Counter
    c = Counter()
    with _open_text(args.dict) as f:
        for line in f:
            try: e = json.loads(line)
            except ValueError: continue
            if e.get("lang_code") != "tl" or e.get("pos") != "noun": continue
            names = set(_cat_names(e.get("categories")))
            for s in e.get("senses", []):
                names |= set(_cat_names(s.get("categories"))); names |= {"topic:" + t for t in s.get("topics", [])}
            c.update(n for n in names if "terms with" not in n and "entries" not in n and "lemmas" not in n and "pages" not in n.lower())
    print("count  category / topic (Tagalog noun entries). Copy useful names into DICT_RULES.")
    for n, k in c.most_common(args.top): print(f"{k:6d}  {n}")

def cmd_discover(args):
    cl = Client(args.cache, args.contact, args.delay)
    out = []
    for host in args.hosts.split(","):
        for pref in args.prefixes.split(","):
            cont = {}
            while True:
                p = dict(action="query", list="allcategories", acprefix=pref, aclimit=500, acprop="size",
                         format="json", formatversion=2, **cont)
                data = cl.get(api(host), p) or {}
                for c in data.get("query", {}).get("allcategories", []):
                    out.append((host, c["category"], c.get("pages", 0), c.get("subcats", 0)))
                if "continue" in data: cont = data["continue"]
                else: break
    out = sorted(set(out), key=lambda r: (r[0], -r[2]))
    with open(args.out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["host", "category", "pages", "subcats"]); w.writerows(out)
    print(f"Wrote {len(out)} categories to {args.out}. Copy useful names into POOLS.")

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    def common(p):
        p.add_argument("--contact", default="unset@example.com", help="email/URL for the User-Agent (Wikimedia policy)")
        p.add_argument("--cache", default=".pool_cache"); p.add_argument("--delay", type=float, default=0.15)
    r = sub.add_parser("run"); common(r)
    r.add_argument("--pools", help="comma-separated pool names (default: all)")
    r.add_argument("--out", default="replacement_pool.xlsx"); r.add_argument("--csv")
    r.add_argument("--max-per-source", type=int, default=600); r.add_argument("--no-views", action="store_true")
    r.add_argument("--workers", type=int, default=4)
    r.add_argument("--dict", help="kaikki.org Tagalog JSONL (or .gz): Tagalog dictionary source for generic-noun pools")
    r.add_argument("--dict-cap", type=int, default=400, help="max dictionary rows per pool (best matches first)")
    r.add_argument("--allow-english-loans", action="store_true", help="keep words Wiktionary marks as borrowed from English")
    r.add_argument("--names-file", help="CSV/TXT  name[,count[,gender]]  to fill the Person name pool (e.g. a PSA/baby-name list)")
    r.add_argument("--targets", default="GSMPH_Cleaned_Terms.xlsx", help="term-bank workbook; adds the Coverage sheet (skipped if missing)")
    r.add_argument("--buffer", type=float, default=0.10, help="excess over tagged unique terms (0.10 = 10%%)")
    r.add_argument("--min-score", type=float, default=3.0,
                   help="Wikipedia candidates scoring below this go to 'Rejected (auto)'; 0 disables. Needs page views (not --no-views)")
    d = sub.add_parser("discover"); common(d)
    d.add_argument("--hosts", default="en.wikipedia.org,tl.wikipedia.org")
    d.add_argument("--prefixes", default="Philippine,Filipino,Tagalog,Cebuano,Visayan")
    d.add_argument("--out", default="categories_found.csv")
    fd = sub.add_parser("fetch-dict"); fd.add_argument("--contact", default="unset@example.com")
    fd.add_argument("--url", default=KAIKKI_URL); fd.add_argument("--out", default="kaikki_tagalog.jsonl")
    dd = sub.add_parser("discover-dict"); dd.add_argument("dict"); dd.add_argument("--top", type=int, default=80)
    t = sub.add_parser("targets"); t.add_argument("--targets", default="GSMPH_Cleaned_Terms.xlsx")
    t.add_argument("--buffer", type=float, default=0.10); t.add_argument("--csv")
    sub.add_parser("list-pools")
    a = ap.parse_args()
    if a.cmd == "list-pools":
        for k, v in POOLS.items():
            kind = ("wiki " if v["sources"] else "") + ("dict " if v.get("dict") else "") + ("curated" if k in CURATED else "")
            print(f"{k:34s} [{kind.strip():16s}] " + "; ".join(s["category"] for s in v["sources"]))
    elif a.cmd == "run": cmd_run(a)
    elif a.cmd == "targets": cmd_targets(a)
    elif a.cmd == "fetch-dict": cmd_fetch_dict(a)
    elif a.cmd == "discover-dict": cmd_discover_dict(a)
    else: cmd_discover(a)

if __name__ == "__main__":
    main()