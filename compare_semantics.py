import pandas as pd
import openpyxl
from sentence_transformers import SentenceTransformer, util

excel_file = 'GSMPH_Cleaned_Terms.xlsx'
target_sheet = 'Lookup Table (Cleaned)'

print("1. Loading SBERT Multilingual model...")
model = SentenceTransformer('paraphrase-multilingual-MiniLM-L12-v2')

print(f"2. Reading sheet '{target_sheet}' from {excel_file}...")
df = pd.read_excel(excel_file, sheet_name=target_sheet, engine='openpyxl')

col_orig = 'Term'
col_loc = 'Previous Replacement' if 'Previous Replacement' in df.columns else 'Replacement Pool'

print(f"3. Comparing semantic similarity between: '{col_orig}' VS '{col_loc}'...")
original_texts = df[col_orig].fillna('').astype(str).tolist()
localized_texts = df[col_loc].fillna('').astype(str).tolist()

embeddings_orig = model.encode(original_texts, batch_size=64, convert_to_tensor=True, show_progress_bar=True)
embeddings_loc = model.encode(localized_texts, batch_size=64, convert_to_tensor=True, show_progress_bar=True)

cosine_scores = util.cos_sim(embeddings_orig, embeddings_loc)
scores = [cosine_scores[i][i].item() for i in range(len(df))]

print(f"4. Directly writing scores into '{excel_file}' without changing anything else...")

wb = openpyxl.load_workbook(excel_file)
ws = wb[target_sheet]

col_score_idx = 15
col_review_idx = 16

ws.cell(row=1, column=col_score_idx, value='semantic_similarity_score')
ws.cell(row=1, column=col_review_idx, value='needs_human_review')

for i, score in enumerate(scores, start=2):
    ws.cell(row=i, column=col_score_idx, value=score)
    ws.cell(row=i, column=col_review_idx, value=(score < 0.70))

wb.save(excel_file)
wb.close()

print(f"\nSUCCESS!\nSemantic similarity scores and human review flags have been added to '{excel_file}' in sheet '{target_sheet}'.")